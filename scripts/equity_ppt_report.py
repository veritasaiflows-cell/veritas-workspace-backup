import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

WORKSPACE = Path(r"C:\Users\Veritas2.0\.openclaw\workspace")
TMP = WORKSPACE / "tmp"
PLAYBOOKS = WORKSPACE / "06. Playbooks"
TMP.mkdir(parents=True, exist_ok=True)
PLAYBOOKS.mkdir(parents=True, exist_ok=True)

BG = RGBColor(245, 241, 232)
CARD = RGBColor(237, 231, 218)
TEXT = RGBColor(17, 17, 17)
MUTED = RGBColor(63, 63, 63)
GREEN = RGBColor(13, 138, 83)
RED = RGBColor(185, 71, 40)


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.34), Inches(11.6), Inches(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = TEXT
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.92), Inches(11.4), Inches(0.4))
        tf2 = sub_box.text_frame
        p2 = tf2.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(11)
        r2.font.color.rgb = MUTED


def add_bullets(slide, items, left, top, width, height, font_size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = item
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(6)


def add_image(slide, path, left, top, width=None, height=None):
    path = Path(path)
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width) if width else None, height=Inches(height) if height else None)


def add_metric_card(slide, title, value, note, left, top, width=3.45, height=1.45, note_color=None):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = CARD
    tx = slide.shapes.add_textbox(Inches(left + 0.18), Inches(top + 0.11), Inches(width - 0.32), Inches(height - 0.2))
    tf = tx.text_frame
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = title
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = TEXT
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = value
    r2.font.size = Pt(20)
    r2.font.bold = True
    r2.font.color.rgb = TEXT
    p3 = tf.add_paragraph()
    r3 = p3.add_run()
    r3.text = note
    r3.font.size = Pt(10)
    r3.font.color.rgb = note_color or MUTED


def build_generic_deck(data, out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    ticker = data["ticker"]
    latest = data["latest"]
    f = data["fundamentals"]
    verdict = data.get("verdict", "WATCH / BENCH")
    verdict_reason = data.get("verdict_reason", "Setup still needs work.")
    safe = ticker.lower().replace(".", "-")

    price_panel = TMP / f"{safe}-price-panel.png"
    history_panel = TMP / f"{safe}-history-panel.png"
    earnings_panel = TMP / f"{safe}-earnings-panel.png"
    segment_panel = TMP / f"{safe}-segment-panel.png"

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, f"{f['company']} ({ticker})", "Visual investment deck | Generated from Veritas report assets")
    add_bullets(slide, [
        f"Current status: {verdict}",
        verdict_reason,
        f"Current price {data['snapshot']['last_close']} vs consensus target {data['snapshot']['consensus_target']}",
        f"Analyst rating: {f['analyst_rating']}"
    ], left=0.85, top=1.58, width=4.8, height=3.1, font_size=18)
    add_image(slide, price_panel, 5.42, 1.28, width=6.95)

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "Executive Summary", "What matters now")
    add_metric_card(slide, "Close", data['snapshot']['last_close'], f"52-wk high ${latest['high_52w']:.2f}", 0.9, 1.45, note_color=RED if latest['close'] < latest['high_52w'] else GREEN)
    add_metric_card(slide, "Consensus target", data['snapshot']['consensus_target'], f"{f['analyst_count']} analysts", 4.58, 1.45, note_color=GREEN if '-' not in data['snapshot']['target_upside'] else RED)
    add_metric_card(slide, "Current stance", verdict, "Action status now", 8.26, 1.45, note_color=RED if "DO NOT" in verdict or "PULLBACK" in verdict else GREEN)
    add_bullets(slide, [
        f"Business quality remains credible in {f['sector']} / {f['industry']}",
        f"Market cap {data['snapshot']['market_cap']} with {data['snapshot']['inst_owned']} institutional ownership",
        f"Forward P/E {data['snapshot']['forward_pe']} and dividend yield {data['snapshot']['dividend_yield']}",
        f"Target upside currently {data['snapshot']['target_upside']}"
    ], left=1.0, top=3.38, width=4.95, height=2.75, font_size=18)
    add_bullets(slide, [
        f"20D / 50D / 200D: {latest['ma20']:.1f} / {latest['ma50']:.1f} / {latest['ma200']:.1f}",
        f"Support / resistance: {latest['recent_support']:.1f} / {latest['recent_resistance']:.1f}",
        "Use the chart to decide timing, not the business story alone",
        "Respect entry quality and opportunity cost"
    ], left=6.55, top=3.38, width=4.8, height=2.75, font_size=18)

    if earnings_panel.exists():
        slide = prs.slides.add_slide(blank)
        set_bg(slide)
        add_title(slide, "Quarter and Guidance", "Company-specific overlay")
        add_image(slide, earnings_panel, 0.82, 1.28, width=11.7)

    if segment_panel.exists():
        slide = prs.slides.add_slide(blank)
        set_bg(slide)
        add_title(slide, "Segment Performance", "Company-specific overlay")
        add_image(slide, segment_panel, 0.82, 1.28, width=11.7)

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "Historical Financial Trend", "Ticker-level revenue and EPS history")
    add_image(slide, history_panel, 0.82, 1.28, width=11.7)
    add_bullets(slide, [
        "Use this slide to check whether the business trend supports the story",
        "A good trend does not eliminate timing risk, but it does separate quality from hype"
    ], left=0.95, top=6.28, width=10.95, height=0.8, font_size=16)

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "Technical and Action Read", "Why the current verdict is what it is")
    add_metric_card(slide, "20D / 50D / 200D", f"{latest['ma20']:.1f} / {latest['ma50']:.1f} / {latest['ma200']:.1f}", "Trend posture snapshot", 0.9, 1.55, note_color=RED if latest['close'] < latest['ma20'] else GREEN)
    add_metric_card(slide, "Support / resistance", f"{latest['recent_support']:.1f} / {latest['recent_resistance']:.1f}", "Recent range context", 4.58, 1.55, note_color=MUTED)
    add_metric_card(slide, "Action", verdict, "Veritas action state", 8.26, 1.55, note_color=RED if "DO NOT" in verdict or "PULLBACK" in verdict else GREEN)
    add_bullets(slide, [
        verdict_reason,
        "Ask whether the chart supports action now or only after a better setup forms",
        "Keep the business case separate from the entry decision"
    ], left=1.0, top=3.45, width=5.0, height=2.2, font_size=18)
    add_bullets(slide, [
        "Upgrade only after cleaner trend structure or better entry quality",
        "Downgrade if the setup deteriorates further or valuation support weakens",
        "Use the deployment sheet for final portfolio action"
    ], left=6.55, top=3.45, width=4.6, height=2.2, font_size=18)

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "Decision Slide", "Bottom line")
    add_metric_card(slide, "Verdict", verdict, "Current action state", 0.9, 1.6, note_color=RED if "DO NOT" in verdict or "PULLBACK" in verdict else GREEN)
    add_metric_card(slide, "Target upside", data['snapshot']['target_upside'], "Vs consensus target", 4.58, 1.6, note_color=GREEN if '-' not in data['snapshot']['target_upside'] else RED)
    add_metric_card(slide, "Key level", f"{latest['recent_support']:.1f}", "Support to respect", 8.26, 1.6, note_color=MUTED)
    add_bullets(slide, [
        f"{ticker} remains in active coverage",
        verdict_reason,
        "Do not let a good story substitute for good timing"
    ], left=1.05, top=3.92, width=4.55, height=1.85, font_size=20)
    add_bullets(slide, [
        "Use this deck as orientation, not as blind execution permission",
        "Final deployment should still respect the written risk and trigger framework"
    ], left=6.95, top=3.92, width=4.35, height=1.85, font_size=18)

    prs.save(out_path)


def main():
    parser = argparse.ArgumentParser(description="Generate a PowerPoint deck from visual equity report assets.")
    parser.add_argument("ticker", help="Ticker symbol, e.g. RTX")
    parser.add_argument("--data-json", default=None, help="Optional JSON data path. Defaults to tmp/<ticker>-visual-report-data.json then analysis-data-v2")
    parser.add_argument("--out", default=None, help="Optional output .pptx path")
    args = parser.parse_args()

    ticker = args.ticker.upper()
    safe = ticker.lower().replace('.', '-')
    default_paths = [
        TMP / f"{safe}-visual-report-data.json",
        TMP / f"{safe}-analysis-data-v2.json",
        TMP / "rtx-analysis-data-v2.json",
    ]
    data_path = Path(args.data_json) if args.data_json else next((p for p in default_paths if p.exists()), None)
    if not data_path or not data_path.exists():
        raise SystemExit(f"No data JSON found for {ticker}. Expected one of: {default_paths}")

    data = load_json(data_path)
    out_path = Path(args.out) if args.out else PLAYBOOKS / f"{ticker} Deck - {data['latest']['date']}.pptx"
    build_generic_deck(data, out_path)

    print(json.dumps({
        "ticker": ticker,
        "data_path": str(data_path),
        "deck_path": str(out_path)
    }, indent=2))


if __name__ == "__main__":
    main()
