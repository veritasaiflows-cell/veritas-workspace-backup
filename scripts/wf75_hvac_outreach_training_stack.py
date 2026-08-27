#!/usr/bin/env python3
"""Render the WF75 HVAC outreach academy training stack.

This packages the already-grounded HVAC private-pilot outreach packet into
interactive and printable training assets. It is internal training only. It
does not approve outreach, payments, channel bindings, customer data
collection, CRM import, credentials, production access, finance/account
actions, or owner approval.
"""
from __future__ import annotations

import argparse
import html
import json
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from market_data_utils import atomic_write_json, atomic_write_text, load_json_artifact


ROOT = Path(__file__).resolve().parents[1]
ACADEMY = ROOT / "training" / "wf75-academy"
TMP = ROOT / "tmp"
SOURCE_PACKET = TMP / "wf75-hvac-prospect-outreach-20260703" / "packet.json"

SOURCE_JSON = ACADEMY / "hvac-prospect-outreach-training-stack-2026-07-03.json"
TRAINING_MD = ACADEMY / "hvac-prospect-outreach-training-2026-07-03.md"
SIM_LAB_MD = ACADEMY / "hvac-prospect-outreach-simulation-lab-2026-07-03.md"
HANDOUT_HTML = ACADEMY / "hvac-prospect-outreach-handout-2026-07-03.html"
HANDOUT_PDF = ACADEMY / "hvac-prospect-outreach-handout-2026-07-03.pdf"
SIM_HTML = ACADEMY / "hvac-prospect-outreach-simulation-deck-2026-07-03.html"
DOCX_BOOK = ACADEMY / "hvac-prospect-outreach-training-book-2026-07-03.docx"
PPTX_DECK = ACADEMY / "hvac-prospect-outreach-activity-deck-2026-07-03.pptx"
MANIFEST = ACADEMY / "hvac-prospect-outreach-training-manifest-2026-07-03.json"
RENDER_PROOF = TMP / "wf75-hvac-outreach-training-stack-render.json"


AUTHORITY_BOUNDARY = {
    "internal_training_only": True,
    "customer_outreach_approved": False,
    "payment_collection_approved": False,
    "customer_data_collection_approved": False,
    "channel_binding_approved": False,
    "crm_import_approved": False,
    "bulk_automation_approved": False,
    "phone_sms_approved": False,
    "credential_or_production_access_approved": False,
    "owner_approval_inferred": False,
}


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return str(path)


def esc(value: Any) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def clean_text(text: str) -> str:
    return "\n".join(line.rstrip() for line in text.splitlines()) + "\n"


def find_browser() -> str | None:
    candidates = [
        shutil.which("msedge"),
        shutil.which("chrome"),
        shutil.which("chromium"),
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return str(candidate)
    return None


def simulations() -> list[dict[str, Any]]:
    return [
        {
            "id": "truth-surface",
            "title": "Locate The Truth Surface",
            "difficulty": "foundation",
            "scenario": "Randall asks where the current outreach packet is and what it authorizes.",
            "task": "Answer in 60 seconds. Name the packet location and authority boundary.",
            "expected": (
                "The packet is in 10. Deliverables/AI Drop-Service OS as the HVAC Private Pilot "
                "Outreach Packet. It prepares outreach but does not authorize sending. Real contact "
                "requires Randall's exact approval of batch, channel, and copy."
            ),
            "fail_triggers": ["QA cleared it for sending", "cannot name file family", "omits owner-gated boundary"],
            "category": "source",
        },
        {
            "id": "offer-plain-english",
            "title": "Explain The Offer Without Hype",
            "difficulty": "foundation",
            "scenario": "A friend asks what Randall is selling.",
            "task": "Explain the AI Workflow Clarity Sprint in plain English.",
            "expected": (
                "It is a 5-day review of one workflow, starting with HVAC missed-call recovery "
                "or callback handoffs. The output is a practical map and improvement checklist. "
                "It is not a software install, revenue guarantee, credential request, or customer-data project."
            ),
            "fail_triggers": ["AI will automate everything", "promises more bookings", "implementation included"],
            "category": "offer",
        },
        {
            "id": "safe-angle",
            "title": "Pick The Safe Outreach Angle",
            "difficulty": "foundation",
            "scenario": "A prospect's public site emphasizes online booking, fast response, and emergency service.",
            "task": "Write one safe outreach angle.",
            "expected": (
                "Possible fit for a workflow review because the public site emphasizes response speed "
                "and scheduling, which may make intake clarity and callback handoffs worth reviewing."
            ),
            "fail_triggers": ["they are losing leads", "their workflow is broken", "they need AI"],
            "category": "copy",
        },
        {
            "id": "unsafe-copy-repair",
            "title": "Repair Unsafe Copy",
            "difficulty": "core",
            "scenario": "Unsafe draft: Your HVAC company is probably losing calls and revenue. Our AI can recover those missed opportunities.",
            "task": "Rewrite it safely.",
            "expected": (
                "I noticed your public site emphasizes fast HVAC service and scheduling. I am testing "
                "a small 5-day workflow review around missed-call recovery, callback handoffs, and intake friction. "
                "It is a diagnostic review, not a promise of more bookings or revenue."
            ),
            "fail_triggers": ["keeps revenue claim", "keeps accusation", "keeps AI can recover"],
            "category": "copy",
        },
        {
            "id": "channel-boundary",
            "title": "Channel Boundary",
            "difficulty": "core",
            "scenario": "Randall says: Let's call all 20 tomorrow.",
            "task": "Respond with the correct boundary.",
            "expected": (
                "Do not call all 20. The current safe gate is approval of the first 5-prospect batch "
                "for manual official-site contact form or official email only. Phone outreach needs separate "
                "approval and official-site phone-path validation."
            ),
            "fail_triggers": ["phone approved", "20 is fine", "bulk automation"],
            "category": "channel",
        },
        {
            "id": "send-details",
            "title": "Prospect Replies: Send Details",
            "difficulty": "core",
            "scenario": "After approved outreach, a prospect replies: Sure, send details.",
            "task": "Draft a safe reply.",
            "expected": (
                "Thanks. I can send the one-page scope. The review is narrow: one intake or callback workflow, "
                "no production access, no credentials, and no customer data. The goal is a practical workflow map "
                "and improvement checklist. If that sounds aligned, a 15-minute call can confirm fit."
            ),
            "fail_triggers": ["sends payment link", "requests CRM access", "offers implementation"],
            "category": "reply",
        },
        {
            "id": "recordings-stop",
            "title": "Prospect Offers Call Recordings",
            "difficulty": "core",
            "scenario": "Prospect says they can send call recordings and customer texts.",
            "task": "Respond safely.",
            "expected": (
                "Not for this first review. Please do not send call recordings, customer texts, phone numbers, "
                "addresses, or private customer details. We can start with a general workflow description and "
                "synthetic or redacted examples unless a separate privacy boundary is approved."
            ),
            "fail_triggers": ["accepts recordings", "asks for transcripts", "uploads customer data"],
            "category": "data",
        },
        {
            "id": "discovery-open",
            "title": "Discovery Call Opening",
            "difficulty": "core",
            "scenario": "You are opening a 15-minute discovery call with an HVAC owner.",
            "task": "Give the first 30 seconds.",
            "expected": (
                "Thanks for taking the time. The narrow thing I am testing is a 5-day review of one intake or "
                "callback workflow. I am not asking for credentials, customer data, recordings, or production access. "
                "Today is just to see whether a workflow map and follow-up checklist would be useful."
            ),
            "fail_triggers": ["starts selling tools", "asks for login", "opens with ROI"],
            "category": "call",
        },
        {
            "id": "servicetitan",
            "title": "Objection: We Already Use ServiceTitan",
            "difficulty": "advanced",
            "scenario": "Prospect says they already use ServiceTitan.",
            "task": "Respond without criticizing their stack.",
            "expected": (
                "That may be fine. This Sprint is not a replacement for ServiceTitan. It looks at the human workflow "
                "around missed calls, callbacks, and handoff status. If that is already clean, the Sprint is probably not needed."
            ),
            "fail_triggers": ["attacks ServiceTitan", "claims AI improves it", "pushes implementation"],
            "category": "objection",
        },
        {
            "id": "revenue-question",
            "title": "Objection: How Much Revenue Will This Add?",
            "difficulty": "advanced",
            "scenario": "Prospect asks how much revenue this will add.",
            "task": "Answer safely.",
            "expected": (
                "I cannot promise revenue or booked-job lift. The first Sprint is a diagnostic. It clarifies where "
                "follow-up, ownership, or handoff status may be unclear. Any later measurement must be defined separately "
                "and treated as evidence, not a guarantee."
            ),
            "fail_triggers": ["gives a percentage", "says it pays for itself", "guaranteed ROI"],
            "category": "objection",
        },
        {
            "id": "scope-classifier",
            "title": "Qualification Decision",
            "difficulty": "advanced",
            "scenario": "A prospect wants text automation, missed-call auto-replies, and after-hours emergency triage.",
            "task": "Classify the prospect.",
            "expected": (
                "Interested but out of first-scope. Text automation and emergency triage require separate approval, "
                "compliance review, consent and opt-out handling, and implementation scope. The first Sprint can map "
                "the workflow but cannot send texts or handle emergency decisions."
            ),
            "fail_triggers": ["accepts text automation", "includes emergency triage", "skips compliance boundary"],
            "category": "qualification",
        },
        {
            "id": "approval-card",
            "title": "Build The Approval Card",
            "difficulty": "advanced",
            "scenario": "Randall wants to proceed with the first batch.",
            "task": "Fill the minimum approval fields.",
            "expected": (
                "Decision: approve manual first-batch outreach. Batch: Olive Air, Instant Heating and Air, Frozen Cactus, "
                "Cold Stinger, Norris Air. Channel: official-site contact form or official website email only. Copy: approved "
                "initial email and one follow-up. Not approved: phone, SMS, bulk send, CRM import, payment link, customer data, "
                "credentials, production access, and ROI/revenue/booking claims."
            ),
            "fail_triggers": ["approval already granted", "omits channel", "omits stop lines"],
            "category": "approval",
        },
    ]


def build_training() -> dict[str, Any]:
    packet = load_json_artifact(SOURCE_PACKET)
    prospects = packet.get("top_5", [])
    return {
        "schema": "veritas.wf75.hvac_outreach_training_stack.v1",
        "generated_at_utc": utc_now(),
        "status": "internal_training_full_stack_ready",
        "source_packet": rel(SOURCE_PACKET),
        "source_packet_status": packet.get("status"),
        "qa_decision": packet.get("qa_decision"),
        "authority_boundary": AUTHORITY_BOUNDARY,
        "source_locations": {
            "outreach_packet": "10. Deliverables/AI Drop-Service OS/AI Workflow Clarity Sprint - HVAC Private Pilot Outreach Packet - 2026-07-03.md",
            "research_scout_findings": "C:/Users/Veritas/.openclaw/workspaces/research-scout/findings/wf75/hvac-prospect-research-20260703.md",
            "qa_redteam_audit": "C:/Users/Veritas/.openclaw/workspaces/qa-redteam/audits/wf75/hvac-prospect-outreach-audit-20260703.md",
            "prospects_csv": "tmp/wf75-hvac-prospect-outreach-20260703/prospects.csv",
            "packet_json": rel(SOURCE_PACKET),
        },
        "offer": {
            "name": "AI Workflow Clarity Sprint",
            "niche": "HVAC",
            "wedge": "missed-call recovery, callback handoffs, and intake friction",
            "plain_english": (
                "A fixed-scope 5-day review of one intake or callback workflow. "
                "It returns a practical map and improvement checklist without system access, "
                "customer data, production changes, or outcome guarantees."
            ),
            "included": [
                "current-state workflow map",
                "handoff and follow-up gap list",
                "safe acknowledgement and escalation sketch",
                "implementation-readiness checklist",
                "manual status-tracking template",
            ],
            "not_included": [
                "guaranteed revenue or booked jobs",
                "live automation build",
                "CRM migration",
                "text automation",
                "call recording setup",
                "emergency dispatch coverage",
                "legal, compliance, or security certification",
            ],
        },
        "first_batch": prospects,
        "workflow": [
            {"phase": "Research Scout", "output": "public-source prospect evidence", "authority": "internal research"},
            {"phase": "Veritas main", "output": "outreach packet and approval card", "authority": "integration and truth gate"},
            {"phase": "QA Red-Team", "output": "claim/privacy/proof audit", "authority": "internal readiness challenge"},
            {"phase": "Randall", "output": "approve/revise exact batch, channel, copy", "authority": "owner-gated external action"},
            {"phase": "Manual test", "output": "official-site contact only if approved", "authority": "bounded external contact"},
            {"phase": "Response handling", "output": "manual tracking and safe replies", "authority": "no payment/data without new approval"},
        ],
        "safe_copy": {
            "initial": (
                "I was looking at {business_name}'s public HVAC service pages and noticed you already emphasize "
                "fast service, scheduling, and customer response. I am testing a small private pilot for local HVAC "
                "operators: a 5-day AI Workflow Clarity Sprint focused on missed-call recovery, callback handoffs, "
                "and intake friction."
            ),
            "follow_up": (
                "Following up once. I reached out because your public site highlights scheduling, service response, "
                "or broad service coverage, and I thought a workflow review might be relevant if intake handoffs are "
                "something you already care about. The Sprint is a fixed-scope review, not a promise of more bookings or revenue."
            ),
            "unsafe_patterns": [
                "you are losing leads",
                "we can increase bookings",
                "AI will fix your missed calls",
                "guaranteed ROI",
                "send call recordings",
                "send CRM access",
            ],
        },
        "rubric": [
            {"category": "offer clarity", "ready_standard": "explains the Sprint in plain English without hype"},
            {"category": "scope control", "ready_standard": "states what is not included"},
            {"category": "channel discipline", "ready_standard": "does not send, call, text, or automate without approval"},
            {"category": "data discipline", "ready_standard": "blocks customer data, credentials, recordings, and production access"},
            {"category": "claim discipline", "ready_standard": "avoids ROI, revenue, booked-job, emergency, legal, compliance, and security claims"},
            {"category": "prospect judgment", "ready_standard": "uses public fit signals cautiously"},
            {"category": "escalation", "ready_standard": "routes out-of-scope requests back to Veritas/Randall"},
        ],
        "simulations": simulations(),
        "readiness_standard": {
            "minimum_score_per_category": 2,
            "required_three_score_categories": ["claim discipline", "data discipline", "channel discipline"],
            "owner_approval_still_required_after_pass": True,
        },
    }


def flow_svg() -> str:
    labels = ["Research", "Draft", "QA", "Approve", "Manual", "Track"]
    colors = ["#0f766e", "#2563eb", "#7c3aed", "#b45309", "#047857", "#374151"]
    parts = [
        '<svg class="flow-svg" viewBox="0 0 960 170" role="img" aria-label="WF75 outreach workflow">',
        '<rect width="960" height="170" rx="16" fill="#f8fafc"/>',
    ]
    for i, (label, color) in enumerate(zip(labels, colors)):
        x = 34 + i * 151
        parts.append(f'<rect x="{x}" y="45" width="118" height="64" rx="10" fill="{color}"/>')
        parts.append(f'<text x="{x + 59}" y="82" text-anchor="middle" font-size="17" font-weight="700" fill="white">{label}</text>')
        if i < len(labels) - 1:
            ax = x + 123
            parts.append(f'<path d="M{ax} 77 L{ax + 31} 77" stroke="#64748b" stroke-width="4"/>')
            parts.append(f'<path d="M{ax + 31} 77 l-10 -8 v16 z" fill="#64748b"/>')
    parts.append('<text x="480" y="138" text-anchor="middle" font-size="15" fill="#475569">Every external step stays owner-gated until Randall approves exact batch, channel, and copy.</text>')
    parts.append("</svg>")
    return "".join(parts)


def gate_svg() -> str:
    return """
<svg class="gate-svg" viewBox="0 0 760 210" role="img" aria-label="Training readiness gates">
  <rect width="760" height="210" rx="18" fill="#fff7ed"/>
  <g font-family="Arial, Helvetica, sans-serif">
    <text x="34" y="42" font-size="22" font-weight="700" fill="#7c2d12">Readiness gates</text>
    <circle cx="70" cy="92" r="24" fill="#dcfce7" stroke="#16a34a" stroke-width="3"/>
    <text x="70" y="100" text-anchor="middle" font-size="24" font-weight="700" fill="#166534">1</text>
    <text x="115" y="98" font-size="18" fill="#1f2937">Explain offer without ROI or AI hype</text>
    <circle cx="70" cy="142" r="24" fill="#e0f2fe" stroke="#0284c7" stroke-width="3"/>
    <text x="70" y="150" text-anchor="middle" font-size="24" font-weight="700" fill="#075985">2</text>
    <text x="115" y="148" font-size="18" fill="#1f2937">Pass data, channel, and claim discipline drills</text>
    <rect x="502" y="65" width="190" height="86" rx="12" fill="#fee2e2" stroke="#dc2626" stroke-width="3"/>
    <text x="597" y="99" text-anchor="middle" font-size="18" font-weight="700" fill="#7f1d1d">No send authority</text>
    <text x="597" y="127" text-anchor="middle" font-size="14" fill="#7f1d1d">until exact approval</text>
  </g>
</svg>
""".strip()


def render_markdown(data: dict[str, Any]) -> str:
    lines = [
        "# WF75 Academy - HVAC Prospect Outreach Training",
        "",
        f"Generated: {data['generated_at_utc']}",
        "Status: internal training only",
        "",
        "## Bottom Line",
        "",
        "This is the full enhanced WF75 Academy module for HVAC prospect outreach. It includes a printable handout, interactive simulation deck, activity deck, cumulative training book, source JSON, and manifest.",
        "",
        "It does not approve outreach, calls, SMS, CRM import, payment links, customer data collection, channel bindings, production access, or owner approval.",
        "",
        "## Source Location",
        "",
    ]
    for label, path in data["source_locations"].items():
        lines.extend([f"- {label}: `{path}`"])
    lines.extend([
        "",
        "## Offer Frame",
        "",
        f"- Offer: {data['offer']['name']}",
        f"- Niche: {data['offer']['niche']}",
        f"- Wedge: {data['offer']['wedge']}",
        f"- Plain English: {data['offer']['plain_english']}",
        "",
        "## First Batch Under Review",
        "",
        "| Rank | Business | Market | Contact path | Angle |",
        "| --- | --- | --- | --- | --- |",
    ])
    for prospect in data["first_batch"]:
        lines.append(
            f"| {prospect.get('rank')} | {prospect.get('business_name')} | {prospect.get('market')} | "
            f"{prospect.get('contact_or_booking_url')} | {prospect.get('outreach_angle')} |"
        )
    lines.extend([
        "",
        "## Training Modules",
        "",
        "1. Read the offer frame and authority boundary.",
        "2. Study the first-batch prospect table.",
        "3. Practice safe copy repair.",
        "4. Run the simulation deck.",
        "5. Fill the approval-card template only if Randall wants to consider sending.",
        "",
        "## Simulations",
        "",
    ])
    for idx, sim in enumerate(data["simulations"], start=1):
        lines.extend([
            f"### Simulation {idx}: {sim['title']}",
            "",
            f"- Difficulty: {sim['difficulty']}",
            f"- Scenario: {sim['scenario']}",
            f"- Task: {sim['task']}",
            f"- Expected answer: {sim['expected']}",
            f"- Fail triggers: {', '.join(sim['fail_triggers'])}",
            "",
        ])
    lines.extend([
        "## Top Recommendations",
        "",
        "1. P0 - Auto-safe now: use the interactive HTML simulation deck before any approval decision. Why: it makes unsafe copy, channel overreach, and data-boundary mistakes visible.",
        "2. P1 - Auto-safe now: use the PDF handout for focused review before rehearsal. Why: the PDF is the stable reference surface.",
        "3. P2 - Owner-gated: approve or revise exact first-batch outreach only after rehearsal. Why: real prospect contact is external action.",
        "",
        "Do:",
        "",
        "- Treat this as internal training and rehearsal.",
        "",
        "Don't:",
        "",
        "- Treat the training stack, QA internal clearance, or a high simulation score as send approval.",
        "",
        "Improvement / prevention:",
        "",
        "- Keep future outreach training generated from source packets with PDF, simulation, manifest, and no-send boundary in the same stack.",
    ])
    return "\n".join(lines) + "\n"


def render_sim_lab_markdown(data: dict[str, Any]) -> str:
    lines = [
        "# WF75 Academy - HVAC Prospect Outreach Simulation Lab",
        "",
        "Generated: 2026-07-03",
        "Status: internal practice only",
        "",
        "Use the interactive HTML deck for active practice. This Markdown file is the static answer key and fallback.",
        "",
    ]
    for idx, sim in enumerate(data["simulations"], start=1):
        lines.extend([
            f"## Simulation {idx} - {sim['title']}",
            "",
            f"Difficulty: {sim['difficulty']}",
            "",
            f"Scenario: {sim['scenario']}",
            "",
            f"Task: {sim['task']}",
            "",
            "Expected answer:",
            "",
            f"```text\n{sim['expected']}\n```",
            "",
            "Fail triggers:",
            "",
        ])
        for trigger in sim["fail_triggers"]:
            lines.append(f"- {trigger}")
        lines.append("")
    lines.extend([
        "## Scoring",
        "",
        "- 0: unsafe or incorrect.",
        "- 1: partially correct but needs major correction.",
        "- 2: mostly correct with minor correction.",
        "- 3: ready for supervised use.",
        "",
        "Minimum readiness: no category below 2, and claim/data/channel discipline must be 3. Real outreach still requires Randall's explicit approval.",
    ])
    return "\n".join(lines) + "\n"


def render_handout_html(data: dict[str, Any]) -> str:
    prospect_cards = "\n".join(
        f"""
        <article class="prospect-card">
          <div class="rank">#{esc(p.get('rank'))}</div>
          <h3>{esc(p.get('business_name'))}</h3>
          <p class="muted">{esc(p.get('market'))} / {esc(p.get('confidence'))} confidence</p>
          <p>{esc(p.get('outreach_angle'))}</p>
          <p class="url">{esc(p.get('contact_or_booking_url'))}</p>
        </article>
        """
        for p in data["first_batch"]
    )
    rubric_rows = "\n".join(
        f"<tr><td>{esc(item['category'])}</td><td>{esc(item['ready_standard'])}</td></tr>"
        for item in data["rubric"]
    )
    sim_rows = "\n".join(
        f"<tr><td>{idx}</td><td>{esc(sim['title'])}</td><td>{esc(sim['difficulty'])}</td><td>{esc(sim['category'])}</td></tr>"
        for idx, sim in enumerate(data["simulations"], start=1)
    )
    return f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>WF75 HVAC Prospect Outreach Training Handout</title>
<style>
  @page {{ size: Letter; margin: 0.55in; }}
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; font-family: Arial, Helvetica, sans-serif; color: #1f2937; background: #f8fafc; }}
  main {{ max-width: 1120px; margin: 0 auto; background: white; }}
  section {{ padding: 34px 42px; border-bottom: 1px solid #e5e7eb; break-inside: avoid; }}
  .hero {{ background: #f1f5f9; border-bottom: 6px solid #0f766e; }}
  .kicker {{ text-transform: uppercase; letter-spacing: .12em; font-size: 12px; font-weight: 700; color: #475569; }}
  h1 {{ font-size: 38px; line-height: 1.05; margin: 10px 0 12px; color: #111827; }}
  h2 {{ font-size: 25px; margin: 0 0 14px; color: #111827; }}
  h3 {{ margin: 0 0 6px; color: #111827; }}
  p, li {{ font-size: 15px; line-height: 1.45; }}
  .lead {{ font-size: 18px; max-width: 920px; }}
  .grid {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 14px; }}
  .two {{ display: grid; grid-template-columns: 1fr 1fr; gap: 16px; }}
  .panel, .prospect-card {{ border: 1px solid #d1d5db; border-radius: 8px; padding: 16px; background: #fff; }}
  .panel.good {{ border-left: 6px solid #16a34a; }}
  .panel.warn {{ border-left: 6px solid #f59e0b; }}
  .panel.stop {{ border-left: 6px solid #dc2626; }}
  .muted {{ color: #64748b; }}
  .rank {{ display: inline-block; padding: 4px 9px; border-radius: 999px; background: #e0f2fe; color: #075985; font-weight: 700; font-size: 13px; margin-bottom: 8px; }}
  .url {{ font-size: 12px; color: #475569; overflow-wrap: anywhere; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
  th, td {{ border: 1px solid #d1d5db; padding: 8px; text-align: left; vertical-align: top; }}
  th {{ background: #f3f4f6; }}
  .flow-svg, .gate-svg {{ width: 100%; max-height: 260px; }}
  .copy {{ font-family: Consolas, monospace; background: #f8fafc; border: 1px solid #cbd5e1; padding: 14px; border-radius: 8px; white-space: pre-wrap; }}
  .footer-note {{ font-size: 12px; color: #64748b; }}
  @media print {{
    body {{ background: white; }}
    main {{ max-width: none; }}
    section {{ page-break-inside: avoid; }}
  }}
</style>
</head>
<body>
<main>
  <section class="hero">
    <div class="kicker">WF75 Academy / Internal Training Only</div>
    <h1>HVAC Prospect Outreach Training</h1>
    <p class="lead">A practical training handout for the AI Workflow Clarity Sprint: how to explain the offer, rehearse outreach, handle objections, and preserve the no-send boundary.</p>
  </section>

  <section>
    <h2>Current Status</h2>
    <div class="grid">
      <div class="panel good"><h3>What exists</h3><p>20 public-source prospects, a first 5 batch, QA-reviewed internal outreach packet, and a simulation lab.</p></div>
      <div class="panel warn"><h3>What this enables</h3><p>Training, rehearsal, owner approval-card prep, and safer manual outreach planning.</p></div>
      <div class="panel stop"><h3>What it does not authorize</h3><p>No outreach, sending, calling, SMS, CRM import, payment links, customer data, or production access.</p></div>
    </div>
  </section>

  <section>
    <h2>Operating Flow</h2>
    {flow_svg()}
  </section>

  <section>
    <h2>Offer Frame</h2>
    <div class="two">
      <div class="panel good">
        <h3>Plain-English Offer</h3>
        <p>{esc(data['offer']['plain_english'])}</p>
      </div>
      <div class="panel stop">
        <h3>Do Not Sell This As</h3>
        <ul>{''.join(f'<li>{esc(item)}</li>' for item in data['offer']['not_included'])}</ul>
      </div>
    </div>
  </section>

  <section>
    <h2>First Batch Under Review</h2>
    <div class="grid">{prospect_cards}</div>
  </section>

  <section>
    <h2>Safe Copy Standard</h2>
    <div class="two">
      <div class="panel good">
        <h3>Rehearsal Copy</h3>
        <div class="copy">{esc(data['safe_copy']['initial'])}</div>
      </div>
      <div class="panel stop">
        <h3>Unsafe Patterns</h3>
        <ul>{''.join(f'<li>{esc(item)}</li>' for item in data['safe_copy']['unsafe_patterns'])}</ul>
      </div>
    </div>
  </section>

  <section>
    <h2>Readiness Gates</h2>
    {gate_svg()}
  </section>

  <section>
    <h2>Simulation Map</h2>
    <table>
      <thead><tr><th>#</th><th>Simulation</th><th>Difficulty</th><th>Category</th></tr></thead>
      <tbody>{sim_rows}</tbody>
    </table>
  </section>

  <section>
    <h2>Scoring Rubric</h2>
    <table>
      <thead><tr><th>Category</th><th>Ready Standard</th></tr></thead>
      <tbody>{rubric_rows}</tbody>
    </table>
    <p class="footer-note">Minimum readiness: no category below 2. Claim, data, and channel discipline must score 3. This still does not approve outreach.</p>
  </section>
</main>
</body>
</html>
"""


def render_simulation_html(data: dict[str, Any]) -> str:
    sim_data = {
        "title": "HVAC Prospect Outreach Simulation Deck",
        "prospects": data["first_batch"],
        "simulations": data["simulations"],
        "authority": data["authority_boundary"],
        "offer": data["offer"],
    }
    payload = json.dumps(sim_data, ensure_ascii=False)
    html_text = """<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>WF75 HVAC Outreach Simulation Deck</title>
<style>
  * { box-sizing: border-box; }
  body { margin: 0; font-family: Arial, Helvetica, sans-serif; color: #172033; background: #f7f8fb; }
  .app { min-height: 100vh; display: grid; grid-template-columns: 300px 1fr; }
  aside { background: #111827; color: #f9fafb; padding: 22px; display: flex; flex-direction: column; gap: 18px; }
  main { padding: 28px; display: grid; gap: 18px; align-content: start; }
  h1, h2, h3 { margin: 0; }
  h1 { font-size: 26px; line-height: 1.1; }
  h2 { font-size: 28px; }
  p { line-height: 1.48; }
  .muted { color: #6b7280; }
  aside .muted { color: #cbd5e1; }
  .tabs, .navlist, .scorebar { display: flex; gap: 8px; flex-wrap: wrap; }
  button { border: 1px solid #cbd5e1; background: #fff; color: #111827; padding: 9px 12px; border-radius: 8px; cursor: pointer; font-weight: 700; }
  button:hover { border-color: #2563eb; }
  button.active { background: #2563eb; color: #fff; border-color: #2563eb; }
  aside button { width: 100%; background: #1f2937; color: #fff; border-color: #374151; text-align: left; }
  aside button.active { background: #0f766e; border-color: #0f766e; }
  .card { background: #fff; border: 1px solid #d1d5db; border-radius: 8px; padding: 18px; }
  .grid { display: grid; grid-template-columns: repeat(3, minmax(0, 1fr)); gap: 14px; }
  .two { display: grid; grid-template-columns: 1fr 1fr; gap: 14px; }
  .badge { display: inline-flex; align-items: center; padding: 4px 8px; border-radius: 999px; background: #e0f2fe; color: #075985; font-size: 12px; font-weight: 700; }
  .stop { border-left: 6px solid #dc2626; }
  .good { border-left: 6px solid #16a34a; }
  .warn { border-left: 6px solid #f59e0b; }
  .answer { display: none; background: #f8fafc; border: 1px dashed #94a3b8; padding: 14px; border-radius: 8px; white-space: pre-wrap; }
  .answer.open { display: block; }
  .meter { height: 12px; border-radius: 999px; background: #374151; overflow: hidden; }
  .meter div { height: 100%; width: 0%; background: #22c55e; transition: width .2s ease; }
  .scenario-title { display: flex; justify-content: space-between; gap: 12px; align-items: start; }
  textarea { width: 100%; min-height: 140px; border: 1px solid #cbd5e1; border-radius: 8px; padding: 12px; font: 15px Arial, sans-serif; }
  label { display: flex; gap: 8px; align-items: start; margin: 8px 0; }
  input[type="checkbox"] { margin-top: 3px; }
  .svg-wrap svg { width: 100%; max-height: 230px; }
  @media (max-width: 900px) {
    .app { grid-template-columns: 1fr; }
    .grid, .two { grid-template-columns: 1fr; }
    aside { position: static; }
  }
</style>
</head>
<body>
<div class="app">
  <aside>
    <div>
      <div class="badge">WF75 Academy</div>
      <h1>HVAC Outreach Simulation</h1>
      <p class="muted">Internal practice only. Passing this deck does not approve sending.</p>
    </div>
    <div>
      <p class="muted">Progress</p>
      <div class="meter"><div id="meterFill"></div></div>
      <p id="progressText" class="muted"></p>
    </div>
    <div class="navlist" id="nav"></div>
    <button id="resetBtn">Reset local scores</button>
  </aside>
  <main>
    <div class="tabs">
      <button data-tab="briefing" class="active">Briefing</button>
      <button data-tab="prospects">Prospects</button>
      <button data-tab="simulator">Simulator</button>
      <button data-tab="approval">Approval Gate</button>
    </div>
    <section id="content"></section>
  </main>
</div>
<script>
const DATA = __DATA__;
const storeKey = "wf75-hvac-outreach-sim-scores-v1";
let tab = "briefing";
let current = 0;
let scores = JSON.parse(localStorage.getItem(storeKey) || "{}");

function saveScores() {
  localStorage.setItem(storeKey, JSON.stringify(scores));
}

function scoreCount() {
  return Object.keys(scores).filter((id) => scores[id] !== undefined).length;
}

function updateProgress() {
  const count = scoreCount();
  const pct = Math.round((count / DATA.simulations.length) * 100);
  document.getElementById("meterFill").style.width = pct + "%";
  document.getElementById("progressText").textContent = count + " of " + DATA.simulations.length + " simulations scored";
}

function renderNav() {
  const nav = document.getElementById("nav");
  nav.innerHTML = "";
  DATA.simulations.forEach((sim, idx) => {
    const btn = document.createElement("button");
    btn.textContent = (idx + 1) + ". " + sim.title + (scores[sim.id] !== undefined ? " [" + scores[sim.id] + "]" : "");
    btn.className = tab === "simulator" && current === idx ? "active" : "";
    btn.onclick = () => { tab = "simulator"; current = idx; render(); };
    nav.appendChild(btn);
  });
}

function briefing() {
  return `
    <div class="card good">
      <h2>${DATA.offer.name}</h2>
      <p>${DATA.offer.plain_english}</p>
    </div>
    <div class="two">
      <div class="card">
        <h3>The wedge</h3>
        <p>${DATA.offer.wedge}</p>
      </div>
      <div class="card stop">
        <h3>Hard boundary</h3>
        <p>No outreach, calls, SMS, CRM import, payment link, customer data, credentials, production access, or owner approval is created by this training.</p>
      </div>
    </div>
    <div class="card svg-wrap">${flowSvg()}</div>
  `;
}

function flowSvg() {
  return `<svg viewBox="0 0 960 170" role="img" aria-label="Outreach workflow">
    <rect width="960" height="170" rx="16" fill="#f8fafc"/>
    ${["Research","Draft","QA","Approve","Manual","Track"].map((label, i) => {
      const colors = ["#0f766e","#2563eb","#7c3aed","#b45309","#047857","#374151"];
      const x = 34 + i * 151;
      const arrow = i < 5 ? `<path d="M${x+123} 77 L${x+154} 77" stroke="#64748b" stroke-width="4"/><path d="M${x+154} 77 l-10 -8 v16 z" fill="#64748b"/>` : "";
      return `<rect x="${x}" y="45" width="118" height="64" rx="10" fill="${colors[i]}"/><text x="${x+59}" y="82" text-anchor="middle" font-size="17" font-weight="700" fill="white">${label}</text>${arrow}`;
    }).join("")}
    <text x="480" y="138" text-anchor="middle" font-size="15" fill="#475569">External contact waits for exact Randall approval.</text>
  </svg>`;
}

function prospects() {
  return `<div class="grid">${DATA.prospects.map((p) => `
    <article class="card">
      <span class="badge">#${p.rank} / ${p.market}</span>
      <h3>${p.business_name}</h3>
      <p>${p.outreach_angle}</p>
      <p class="muted">${p.contact_or_booking_url}</p>
    </article>
  `).join("")}</div>`;
}

function simulator() {
  const sim = DATA.simulations[current];
  const answerOpen = document.body.dataset.answerOpen === sim.id;
  return `
    <div class="card">
      <div class="scenario-title">
        <div>
          <span class="badge">${sim.difficulty} / ${sim.category}</span>
          <h2>${current + 1}. ${sim.title}</h2>
        </div>
        <span class="badge">Score: ${scores[sim.id] ?? "not scored"}</span>
      </div>
      <p><strong>Scenario:</strong> ${sim.scenario}</p>
      <p><strong>Task:</strong> ${sim.task}</p>
      <textarea placeholder="Write or rehearse your answer here. This stays local in the browser view and is not saved."></textarea>
      <div class="scorebar">
        <button onclick="setScore('${sim.id}',0)">0 Unsafe</button>
        <button onclick="setScore('${sim.id}',1)">1 Needs work</button>
        <button onclick="setScore('${sim.id}',2)">2 Usable</button>
        <button onclick="setScore('${sim.id}',3)">3 Ready</button>
        <button onclick="toggleAnswer('${sim.id}')">${answerOpen ? "Hide" : "Reveal"} standard</button>
      </div>
      <div class="answer ${answerOpen ? "open" : ""}">
        <strong>Expected answer:</strong>
        <p>${sim.expected}</p>
        <strong>Fail triggers:</strong>
        <ul>${sim.fail_triggers.map((x) => `<li>${x}</li>`).join("")}</ul>
      </div>
    </div>
  `;
}

function approval() {
  const checks = [
    "First 5 prospects are exact and approved",
    "Channel is official-site contact form or official email only",
    "Copy version is exact and approved",
    "No phone, SMS, bulk send, CRM import, or payment link",
    "No customer data, credentials, recordings, or production access",
    "No ROI, revenue, booking, emergency, legal, compliance, or security claims",
    "Manual tracking path is ready"
  ];
  return `
    <div class="card warn">
      <h2>Approval Gate</h2>
      <p>Use this checklist before any real external contact. Checking boxes here does not approve outreach; it prepares the approval decision.</p>
      ${checks.map((c, i) => `<label><input type="checkbox" data-approval="${i}"> ${c}</label>`).join("")}
    </div>
    <div class="card stop">
      <h3>Final sentence before sending</h3>
      <p>Randall approves exact first-batch manual official-site contact using the approved copy: yes or no.</p>
    </div>
  `;
}

function render() {
  document.querySelectorAll("[data-tab]").forEach((btn) => {
    btn.classList.toggle("active", btn.dataset.tab === tab);
  });
  const content = document.getElementById("content");
  if (tab === "briefing") content.innerHTML = briefing();
  if (tab === "prospects") content.innerHTML = prospects();
  if (tab === "simulator") content.innerHTML = simulator();
  if (tab === "approval") content.innerHTML = approval();
  renderNav();
  updateProgress();
}

function setScore(id, value) {
  scores[id] = value;
  saveScores();
  render();
}

function toggleAnswer(id) {
  document.body.dataset.answerOpen = document.body.dataset.answerOpen === id ? "" : id;
  render();
}

document.querySelectorAll("[data-tab]").forEach((btn) => {
  btn.onclick = () => { tab = btn.dataset.tab; render(); };
});
document.getElementById("resetBtn").onclick = () => {
  scores = {};
  saveScores();
  document.body.dataset.answerOpen = "";
  render();
};
document.addEventListener("keydown", (event) => {
  if (tab !== "simulator") return;
  if (event.key === "ArrowRight" && current < DATA.simulations.length - 1) { current++; render(); }
  if (event.key === "ArrowLeft" && current > 0) { current--; render(); }
});
render();
</script>
</body>
</html>
"""
    return html_text.replace("__DATA__", payload)


def render_pdf(html_path: Path, pdf_path: Path) -> dict[str, Any]:
    browser = find_browser()
    if not browser:
        return {"status": "pdf_unavailable", "reason": "no Edge/Chrome/Chromium renderer found"}
    cmd = [
        browser,
        "--headless=new",
        "--disable-gpu",
        "--no-first-run",
        "--disable-extensions",
        f"--print-to-pdf={str(pdf_path.resolve())}",
        html_path.resolve().as_uri(),
    ]
    completed = subprocess.run(cmd, cwd=str(ROOT), text=True, capture_output=True, timeout=120)
    ok = completed.returncode == 0 and pdf_path.exists() and pdf_path.stat().st_size > 0
    return {
        "status": "ok" if ok else "pdf_failed",
        "browser": browser,
        "returncode": completed.returncode,
        "stdout_tail": completed.stdout[-500:],
        "stderr_tail": completed.stderr[-500:],
        "pdf_size_bytes": pdf_path.stat().st_size if pdf_path.exists() else 0,
    }


def render_docx(data: dict[str, Any]) -> dict[str, Any]:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover - environment dependent
        return {"status": "docx_unavailable", "reason": str(exc)}

    doc = Document()
    doc.add_heading("WF75 HVAC Prospect Outreach Training Book", 0)
    doc.add_paragraph("Internal training only. This book does not approve outreach, payment, customer data collection, or external delivery.")
    doc.add_heading("Offer Frame", level=1)
    doc.add_paragraph(data["offer"]["plain_english"])
    doc.add_heading("First Batch", level=1)
    for prospect in data["first_batch"]:
        doc.add_paragraph(
            f"{prospect.get('rank')}. {prospect.get('business_name')} - {prospect.get('market')}: {prospect.get('outreach_angle')}",
            style="List Bullet",
        )
    doc.add_heading("Simulations", level=1)
    for idx, sim in enumerate(data["simulations"], start=1):
        doc.add_heading(f"{idx}. {sim['title']}", level=2)
        doc.add_paragraph(f"Scenario: {sim['scenario']}")
        doc.add_paragraph(f"Task: {sim['task']}")
        doc.add_paragraph(f"Expected: {sim['expected']}")
        doc.add_paragraph("Fail triggers:")
        for trigger in sim["fail_triggers"]:
            doc.add_paragraph(trigger, style="List Bullet")
    doc.add_heading("Boundary", level=1)
    for key, value in data["authority_boundary"].items():
        doc.add_paragraph(f"{key}: {value}", style="List Bullet")
    doc.save(DOCX_BOOK)
    return {"status": "ok", "docx_size_bytes": DOCX_BOOK.stat().st_size}


def render_pptx(data: dict[str, Any]) -> dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
    except Exception as exc:  # pragma: no cover - environment dependent
        return {"status": "pptx_unavailable", "reason": str(exc)}

    prs = Presentation()
    blank = prs.slide_layouts[6]
    ink = RGBColor(31, 41, 55)
    teal = RGBColor(15, 118, 110)

    def add_title(slide: Any, title: str, subtitle: str = "") -> None:
        box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.1), Inches(0.8))
        frame = box.text_frame
        frame.clear()
        p = frame.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.size = Pt(29)
        run.font.bold = True
        run.font.color.rgb = ink
        if subtitle:
            p2 = frame.add_paragraph()
            r2 = p2.add_run()
            r2.text = subtitle
            r2.font.size = Pt(13)
            r2.font.color.rgb = teal

    def add_box(slide: Any, title: str, body: str, left: float, top: float, width: float, height: float) -> None:
        shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
        shape.line.color.rgb = RGBColor(203, 213, 225)
        frame = shape.text_frame
        frame.clear()
        p = frame.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = ink
        p2 = frame.add_paragraph()
        r2 = p2.add_run()
        r2.text = body
        r2.font.size = Pt(12)
        r2.font.color.rgb = ink

    slide = prs.slides.add_slide(blank)
    add_title(slide, "HVAC Prospect Outreach Training", "WF75 Academy / internal training only")
    add_box(slide, "Status", "Full-stack training: PDF, HTML, interactive simulation, book, deck, manifest.", 0.7, 1.5, 5.8, 1.5)
    add_box(slide, "Hard boundary", "No outreach, calls, SMS, CRM import, payments, customer data, credentials, or owner approval inference.", 6.8, 1.5, 5.8, 1.5)
    add_box(slide, "Offer", data["offer"]["plain_english"], 0.7, 3.4, 11.9, 1.8)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "First Batch Under Review", "Manual official-site contact only if Randall approves")
    for idx, prospect in enumerate(data["first_batch"][:5]):
        left = 0.7 + (idx % 2) * 6.05
        top = 1.3 + (idx // 2) * 1.55
        add_box(slide, f"{prospect.get('rank')}. {prospect.get('business_name')}", prospect.get("outreach_angle", ""), left, top, 5.7, 1.25)

    for idx, sim in enumerate(data["simulations"], start=1):
        slide = prs.slides.add_slide(blank)
        add_title(slide, f"Simulation {idx}: {sim['title']}", f"{sim['difficulty']} / {sim['category']}")
        add_box(slide, "Scenario", sim["scenario"], 0.7, 1.35, 11.9, 1.05)
        add_box(slide, "Task", sim["task"], 0.7, 2.7, 5.75, 1.3)
        add_box(slide, "Expected", sim["expected"], 6.8, 2.7, 5.8, 2.2)
        add_box(slide, "Fail triggers", "\n".join(f"- {x}" for x in sim["fail_triggers"]), 0.7, 4.4, 5.75, 1.6)

    prs.save(PPTX_DECK)
    return {"status": "ok", "pptx_size_bytes": PPTX_DECK.stat().st_size}


def validate_outputs(data: dict[str, Any], manifest: dict[str, Any]) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    required_paths = [
        SOURCE_JSON,
        TRAINING_MD,
        SIM_LAB_MD,
        HANDOUT_HTML,
        HANDOUT_PDF,
        SIM_HTML,
        MANIFEST,
    ]
    for path in required_paths:
        if not path.exists() or path.stat().st_size <= 0:
            errors.append(f"missing_or_empty:{rel(path)}")
    if len(data.get("simulations", [])) < 12:
        errors.append("simulation_count_below_12")
    if len(data.get("first_batch", [])) < 5:
        errors.append("first_batch_below_5")
    if "localStorage" not in SIM_HTML.read_text(encoding="utf-8", errors="ignore"):
        errors.append("interactive_progress_storage_missing")
    if "No outreach" not in HANDOUT_HTML.read_text(encoding="utf-8", errors="ignore"):
        errors.append("handout_boundary_missing")
    if manifest.get("pdf_result", {}).get("status") != "ok":
        warnings.append(f"pdf_result:{manifest.get('pdf_result', {}).get('status')}")
    if manifest.get("docx_result", {}).get("status") != "ok":
        warnings.append(f"docx_result:{manifest.get('docx_result', {}).get('status')}")
    if manifest.get("pptx_result", {}).get("status") != "ok":
        warnings.append(f"pptx_result:{manifest.get('pptx_result', {}).get('status')}")
    return {"status": "ok" if not errors else "blocked", "errors": errors, "warnings": warnings}


def write_stack() -> dict[str, Any]:
    ACADEMY.mkdir(parents=True, exist_ok=True)
    data = build_training()
    atomic_write_json(SOURCE_JSON, data)
    atomic_write_text(TRAINING_MD, clean_text(render_markdown(data)))
    atomic_write_text(SIM_LAB_MD, clean_text(render_sim_lab_markdown(data)))
    atomic_write_text(HANDOUT_HTML, clean_text(render_handout_html(data)))
    atomic_write_text(SIM_HTML, clean_text(render_simulation_html(data)))
    pdf_result = render_pdf(HANDOUT_HTML, HANDOUT_PDF)
    docx_result = render_docx(data)
    pptx_result = render_pptx(data)
    manifest = {
        "schema": "veritas.wf75.hvac_outreach_training_manifest.v1",
        "generated_at_utc": utc_now(),
        "status": "pending",
        "outputs": {
            "source_json": rel(SOURCE_JSON),
            "markdown_training": rel(TRAINING_MD),
            "markdown_simulation_lab": rel(SIM_LAB_MD),
            "html_handout": rel(HANDOUT_HTML),
            "pdf_handout": rel(HANDOUT_PDF),
            "html_interactive_simulation": rel(SIM_HTML),
            "docx_book": rel(DOCX_BOOK),
            "pptx_activity_deck": rel(PPTX_DECK),
            "manifest": rel(MANIFEST),
        },
        "source_packet": rel(SOURCE_PACKET),
        "authority_boundary": AUTHORITY_BOUNDARY,
        "simulation_count": len(data["simulations"]),
        "prospect_count": len(data["first_batch"]),
        "pdf_result": pdf_result,
        "docx_result": docx_result,
        "pptx_result": pptx_result,
        "validation": {"status": "pending", "errors": [], "warnings": []},
    }
    validation = validate_outputs(data, manifest)
    manifest["validation"] = validation
    manifest["status"] = "ready" if validation["status"] == "ok" else "blocked"
    atomic_write_json(MANIFEST, manifest)
    atomic_write_json(RENDER_PROOF, {"status": manifest["status"], "manifest": rel(MANIFEST), "validation": validation, "outputs": manifest["outputs"]})
    return manifest


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Render WF75 HVAC outreach academy training stack.")
    parser.add_argument("--write", action="store_true")
    parser.add_argument("--validate", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not SOURCE_PACKET.exists():
        print(json.dumps({"status": "blocked", "error": f"missing source packet: {rel(SOURCE_PACKET)}"}, indent=2))
        return 2
    manifest = write_stack() if args.write else {"status": "dry_run", "validation": {"status": "ok", "errors": [], "warnings": []}}
    print(json.dumps({
        "status": manifest["status"],
        "outputs": manifest.get("outputs"),
        "simulation_count": manifest.get("simulation_count"),
        "prospect_count": manifest.get("prospect_count"),
        "validation": manifest.get("validation"),
        "pdf_result": manifest.get("pdf_result"),
        "docx_result": manifest.get("docx_result"),
        "pptx_result": manifest.get("pptx_result"),
    }, indent=2))
    if args.validate and manifest.get("validation", {}).get("status") != "ok":
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
