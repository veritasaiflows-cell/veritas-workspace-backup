#!/usr/bin/env python3
"""Build the WF75 Academy / Training Desk packet.

This is Randall training and practice infrastructure for the generic
intelligence / SMB workflow automation pivot. It does not authorize outreach,
customer data, credential access, customer-system implementation, external
delivery, subscriptions, spending, public launch, or revenue/ROI claims.
"""
from __future__ import annotations

import argparse
import html
import json
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from market_data_utils import atomic_write_json, atomic_write_text


ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / "tmp"
TRAINING = ROOT / "training"
ACADEMY = TRAINING / "wf75-academy"
OUT_JSON = TMP / "wf75-training-desk-current.json"
OUT_MD = TMP / "wf75-training-desk-current.md"
TRAINING_JSON = ACADEMY / "wf75-academy-current.json"
TRAINING_MD = ACADEMY / "wf75-academy-current.md"
TRAINING_HTML = ACADEMY / "wf75-academy-handout.html"
TRAINING_PDF = ACADEMY / "wf75-academy-handout.pdf"
TRAINING_DOCX = ACADEMY / "wf75-academy-book.docx"
TRAINING_PPTX = ACADEMY / "wf75-academy-activity-deck.pptx"
TRAINING_SIM_HTML = ACADEMY / "wf75-academy-simulation-deck.html"
TRAINING_MANIFEST = ACADEMY / "wf75-academy-manifest.json"

AUTHORITY_FALSE = {
    "real_customer_data_allowed": False,
    "customer_outreach_allowed": False,
    "external_delivery_allowed": False,
    "credential_access_allowed": False,
    "customer_system_implementation_allowed": False,
    "public_launch_allowed": False,
    "spend_or_subscription_allowed": False,
    "guaranteed_revenue_or_roi_claim_allowed": False,
    "legal_compliance_security_readiness_claim_allowed": False,
    "owner_approval_inferred": False,
}


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return path.as_posix()


def esc(value: Any) -> str:
    return html.escape("" if value is None else str(value))


def find_browser() -> str | None:
    candidates = [
        shutil.which("msedge"),
        shutil.which("chrome"),
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    return None


def build_packet() -> dict[str, Any]:
    generated_at = utc_now()
    daily_sessions = [
        {
            "day": 1,
            "title": "SMB Service-Led Offer Basics",
            "duration_minutes": 20,
            "learn": [
                "what the Lead Rescue / Workflow Clarity Sprint is",
                "what it is not",
                "why manual packets come before full SaaS",
            ],
            "practice": "Explain the offer in 90 seconds without saying implementation, guaranteed revenue, or AI magic.",
            "fake_scenario": "Desert Peak HVAC lead follow-up packet",
            "acceptance": "Randall can state scope, exclusions, and next step clearly.",
        },
        {
            "day": 2,
            "title": "Workflow Diagnosis: Trigger To Outcome",
            "duration_minutes": 25,
            "learn": [
                "trigger, handoff, wait state, owner, outcome",
                "where manual workflows break",
                "how to avoid recommending tools too early",
            ],
            "practice": "Map one fake contractor admin workflow from first customer request to invoice.",
            "fake_scenario": "small contractor job-control packet",
            "acceptance": "Randall can identify trigger, bottleneck, owner, and next action.",
        },
        {
            "day": 3,
            "title": "Zapier Fundamentals",
            "duration_minutes": 25,
            "learn": [
                "trigger -> action -> filter -> path -> log",
                "when Zapier is enough",
                "why task volume and permissions matter",
            ],
            "practice": "Design a fake Zap: website form -> owner alert -> Google Sheet row.",
            "fake_scenario": "cleaning company website lead intake",
            "acceptance": "Randall can explain a Zap without touching real accounts.",
        },
        {
            "day": 4,
            "title": "Make Fundamentals",
            "duration_minutes": 25,
            "learn": [
                "scenario, module, route, filter, bundle",
                "when Make is better than Zapier",
                "duplication and route-testing risks",
            ],
            "practice": "Design a fake Make scenario that routes med-spa inquiries by treatment interest.",
            "fake_scenario": "med spa inquiry routing",
            "acceptance": "Randall can explain one Make scenario and its failure mode.",
        },
        {
            "day": 5,
            "title": "Simple CRM / Sheet Control Plane",
            "duration_minutes": 20,
            "learn": [
                "why Google Sheets/Airtable can be enough at first",
                "status labels and owner fields",
                "what belongs in a lead/job tracking table",
            ],
            "practice": "Create fake columns for lead status, next touch, priority, source, owner, and due date.",
            "fake_scenario": "auto repair quote follow-up list",
            "acceptance": "Randall can explain a simple tracking table before recommending CRM.",
        },
        {
            "day": 6,
            "title": "Phone/SMS Workflow Tools",
            "duration_minutes": 25,
            "learn": [
                "OpenPhone vs Twilio mental model",
                "missed-call text-back concepts",
                "SMS consent and opt-out sensitivity",
            ],
            "practice": "Write a safe internal recommendation for missed-call callback priority without sending texts.",
            "fake_scenario": "plumbing missed-call queue",
            "acceptance": "Randall can separate tool recommendation from message sending.",
        },
        {
            "day": 7,
            "title": "Customer Packet QA",
            "duration_minutes": 25,
            "learn": [
                "private-data leakage checks",
                "credential and implementation stop lines",
                "ROI/legal/compliance claim blocks",
            ],
            "practice": "Review one fake customer preview and mark unsafe lines.",
            "fake_scenario": "Lead Rescue customer preview",
            "acceptance": "Randall catches unsafe claims before delivery.",
        },
    ]
    lesson_modules = [
        {
            "day": 1,
            "title": "SMB Service-Led Offer Basics",
            "teaching_objective": "Randall can explain the SMB Workflow Clarity service in plain English, including scope, exclusions, and the next step.",
            "plain_english_lesson": [
                "Small businesses often do not need a new SaaS platform first. They need someone to show them where the current workflow is leaking attention, leads, follow-up, appointments, quotes, or cash.",
                "Our first service is a manual Workflow Clarity packet. We inspect a sanitized or fake scenario, map what is happening, identify where work stalls, and produce a practical owner-facing packet.",
                "The packet can include a daily attention list, lead status cleanup, callback priority, suggested scripts, simple tracking fields, and a tool-stack recommendation.",
                "We do not promise revenue. We do not touch real customer systems yet. We do not ask for credentials. We do not send messages for them. We start with diagnosis and a manual packet because that is faster, safer, and easier to sell than a full SaaS build.",
            ],
            "key_terms": [
                {"term": "Service-led SaaS", "definition": "A manual service that proves the workflow and customer value before software automation is fully built."},
                {"term": "Lead Rescue", "definition": "Finding and prioritizing missed or stalled inquiries so the owner knows who needs attention first."},
                {"term": "Workflow Clarity packet", "definition": "A structured deliverable showing the workflow, bottlenecks, next actions, scripts, and tool recommendations."},
                {"term": "Stop line", "definition": "A boundary we do not cross without explicit approval, such as real customer data, credentials, outreach, or implementation."},
            ],
            "worked_example": {
                "scenario": "Desert Peak HVAC misses 12 calls in a week. Some callers leave voicemails, some submit forms, and quotes are not followed up consistently.",
                "bad_pitch": "We use AI to automate your whole business and recover lost revenue fast.",
                "why_bad": "It overpromises revenue, implies implementation, and sounds like hype.",
                "good_pitch": "We help you see where calls, forms, quotes, and follow-ups are slipping through the cracks. First we create a clear manual packet showing who needs attention, what status each lead is in, what message or callback should happen next, and which simple tools could support the workflow. We do not need your credentials or customer-system access for the first review.",
            },
            "practice_steps": [
                "State the customer pain in one sentence.",
                "State what we produce in one sentence.",
                "State what we do not do yet.",
                "State the next step: a sanitized intake or fake-scenario review.",
            ],
            "self_check": [
                "Did I avoid promising revenue?",
                "Did I avoid saying we will implement everything?",
                "Did I avoid asking for credentials?",
                "Did I explain the manual packet clearly?",
                "Did I give a concrete next step?",
            ],
            "answer_key": "A strong answer says we diagnose workflow leakage and produce a manual owner packet first. It names missed calls/forms/quotes/follow-ups, gives a practical deliverable, preserves boundaries, and asks for a sanitized intake or fake-scenario review as the next step.",
        },
        {
            "day": 2,
            "title": "Workflow Diagnosis: Trigger To Outcome",
            "teaching_objective": "Randall can map a workflow using trigger, handoff, wait state, owner, and outcome.",
            "plain_english_lesson": [
                "A workflow is not just a task list. It is the path from something happening to a result being achieved.",
                "Most SMB workflow problems happen at handoffs and wait states: nobody owns the next touch, the status is unclear, or the same work is tracked in too many places.",
                "Before recommending software, map what triggers the work, who owns each step, where it waits, and what outcome the business actually wants.",
            ],
            "key_terms": [
                {"term": "Trigger", "definition": "The event that starts the workflow."},
                {"term": "Handoff", "definition": "The moment responsibility moves from one person, system, or step to another."},
                {"term": "Wait state", "definition": "Where the workflow pauses until someone acts or information arrives."},
                {"term": "Outcome", "definition": "The business result the workflow exists to produce."},
            ],
            "worked_example": {
                "scenario": "A contractor receives a job request, visits the site, sends a quote, waits for approval, schedules materials, completes work, and sends an invoice.",
                "bad_pitch": "You need a CRM.",
                "why_bad": "It jumps to a tool before identifying the real bottleneck.",
                "good_pitch": "The issue appears to be quote follow-up and invoice timing. We should map owner, due date, and next touch before recommending a tool.",
            },
            "practice_steps": [
                "Name the trigger.",
                "List each handoff.",
                "Mark each wait state.",
                "Assign an owner to the next touch.",
                "Define the outcome.",
            ],
            "self_check": [
                "Did I identify the trigger?",
                "Did I find at least one wait state?",
                "Did I assign ownership?",
                "Did I avoid tool-first thinking?",
            ],
            "answer_key": "A strong workflow diagnosis names the start event, handoffs, wait states, owner, and desired outcome before proposing automation.",
        },
    ]
    return {
        "schema": "veritas.wf75_training_desk_current.v1",
        "generated_at_utc": generated_at,
        "status": "active_training_lane",
        "department_name": "WF75 Academy / Training Desk",
        "purpose": "Train Randall to understand, sell, review, and manually deliver SMB workflow automation packets before fake-scenario practice and later real-client pilots.",
        "current_training_posture": {
            "academy_live": True,
            "mode": "daily_micro_learning_plus_fake_scenario_practice",
            "recommended_daily_time": "20-30 minutes",
            "weekly_review_time": "45-60 minutes",
            "first_skill_goal": "Randall can explain the service, diagnose one workflow, understand basic Zapier/Make patterns, and QA a packet without crossing stop lines.",
        },
        "tool_curriculum": [
            {
                "tool": "Zapier",
                "learn_for": "simple app-to-app automations",
                "starter_pattern": "form submission -> owner alert -> tracking row",
                "do_not_do_yet": "connect real client accounts or send external messages",
            },
            {
                "tool": "Make",
                "learn_for": "branching workflows and data routing",
                "starter_pattern": "inquiry -> classify -> route -> log",
                "do_not_do_yet": "build production scenarios for clients",
            },
            {
                "tool": "n8n",
                "learn_for": "more technical/self-hosted automation concepts",
                "starter_pattern": "webhook -> validate -> transform -> queue",
                "do_not_do_yet": "self-host or maintain client infrastructure",
            },
            {
                "tool": "Google Sheets / Airtable",
                "learn_for": "simple operating tables and status boards",
                "starter_pattern": "lead/job row -> status -> owner -> next touch",
                "do_not_do_yet": "store sensitive real data casually",
            },
            {
                "tool": "OpenPhone / Twilio",
                "learn_for": "phone/SMS workflow concepts",
                "starter_pattern": "missed call -> internal alert -> manual callback queue",
                "do_not_do_yet": "send real customer messages or handle phone-number migration",
            },
        ],
        "daily_micro_learning_plan": daily_sessions,
        "lesson_modules": lesson_modules,
        "fake_scenario_training_bench": [
            "HVAC lead rescue",
            "med spa inquiry routing",
            "contractor job-control and invoice follow-up",
            "auto repair quote follow-up",
            "cleaning company booking workflow",
            "small law firm intake triage",
            "dental recall and appointment confirmation",
        ],
        "repeatable_session_format": [
            "5 minutes: concept in plain English",
            "10 minutes: inspect one fake scenario",
            "10 minutes: design the workflow or automation shape",
            "5 minutes: QA stop-line review",
            "optional 5 minutes: Randall explains it back",
        ],
        "randall_readiness_levels": [
            {
                "level": 1,
                "name": "Understand",
                "description": "Can explain what the service does and does not do.",
            },
            {
                "level": 2,
                "name": "Diagnose",
                "description": "Can map a fake workflow and identify bottlenecks.",
            },
            {
                "level": 3,
                "name": "Recommend",
                "description": "Can name a simple tool pattern without overbuilding.",
            },
            {
                "level": 4,
                "name": "QA",
                "description": "Can catch unsafe claims, data leakage, credential requests, and implementation drift.",
            },
            {
                "level": 5,
                "name": "Pilot Ready",
                "description": "Can review a manual service packet and discuss it with a real prospect after explicit approval gates are cleared.",
            },
        ],
        "stop_lines": list(AUTHORITY_FALSE.keys()),
        "authority_boundary": AUTHORITY_FALSE,
        "validation": {"status": "ok", "errors": [], "warnings": []},
    }


def render_markdown(packet: dict[str, Any]) -> str:
    lines = [
        "# WF75 Academy / Training Desk",
        "",
        f"Status: {packet['status']}",
        f"Generated: {packet['generated_at_utc']}",
        "",
        "## Purpose",
        packet["purpose"],
        "",
        "## Daily Format",
    ]
    lines.extend(f"- {item}" for item in packet["repeatable_session_format"])
    lines.extend(["", "## First 7 Micro-Learning Sessions"])
    for session in packet["daily_micro_learning_plan"]:
        lines.extend([
            "",
            f"### Day {session['day']} - {session['title']}",
            f"Duration: {session['duration_minutes']} minutes",
            "",
            "Learn:",
            *[f"- {item}" for item in session["learn"]],
            "",
            f"Practice: {session['practice']}",
            f"Fake scenario: {session['fake_scenario']}",
            f"Acceptance: {session['acceptance']}",
        ])
    lines.extend(["", "## Actual Lesson Modules"])
    for lesson in packet.get("lesson_modules", []):
        lines.extend([
            "",
            f"### Day {lesson['day']} - {lesson['title']}",
            "",
            f"Objective: {lesson['teaching_objective']}",
            "",
            "Lesson:",
            *[f"- {item}" for item in lesson["plain_english_lesson"]],
            "",
            "Key terms:",
        ])
        lines.extend(f"- {item['term']}: {item['definition']}" for item in lesson["key_terms"])
        example = lesson["worked_example"]
        lines.extend([
            "",
            "Worked example:",
            f"- Scenario: {example['scenario']}",
            f"- Bad pitch: {example['bad_pitch']}",
            f"- Why it is bad: {example['why_bad']}",
            f"- Better pitch: {example['good_pitch']}",
            "",
            "Practice steps:",
            *[f"- {item}" for item in lesson["practice_steps"]],
            "",
            "Self-check:",
            *[f"- {item}" for item in lesson["self_check"]],
            "",
            f"Answer key: {lesson['answer_key']}",
        ])
    lines.extend(["", "## Tool Curriculum"])
    for tool in packet["tool_curriculum"]:
        lines.extend([
            "",
            f"### {tool['tool']}",
            f"- Learn for: {tool['learn_for']}",
            f"- Starter pattern: {tool['starter_pattern']}",
            f"- Do not do yet: {tool['do_not_do_yet']}",
        ])
    lines.extend(["", "## Stop Lines"])
    lines.extend(f"- {item}" for item in packet["authority_boundary"])
    return "\n".join(lines) + "\n"


def render_html_handout(packet: dict[str, Any]) -> str:
    sessions = packet["daily_micro_learning_plan"]
    tools = packet["tool_curriculum"]
    lessons = packet.get("lesson_modules", [])
    session_cards = []
    for session in sessions:
        learn = "".join(f"<li>{esc(item)}</li>" for item in session["learn"])
        session_cards.append(
            f"""
            <section class="card">
              <div class="eyebrow">Day {esc(session['day'])} / {esc(session['duration_minutes'])} minutes</div>
              <h2>{esc(session['title'])}</h2>
              <h3>Learn</h3>
              <ul>{learn}</ul>
              <h3>Activity</h3>
              <p>{esc(session['practice'])}</p>
              <h3>Simulation</h3>
              <p>{esc(session['fake_scenario'])}</p>
              <h3>Acceptance</h3>
              <p>{esc(session['acceptance'])}</p>
            </section>
            """
        )
    tool_rows = "".join(
        f"<tr><td>{esc(tool['tool'])}</td><td>{esc(tool['learn_for'])}</td><td>{esc(tool['starter_pattern'])}</td><td>{esc(tool['do_not_do_yet'])}</td></tr>"
        for tool in tools
    )
    lesson_sections = []
    for lesson in lessons:
        terms = "".join(f"<li><strong>{esc(item['term'])}:</strong> {esc(item['definition'])}</li>" for item in lesson["key_terms"])
        body = "".join(f"<p>{esc(item)}</p>" for item in lesson["plain_english_lesson"])
        practice = "".join(f"<li>{esc(item)}</li>" for item in lesson["practice_steps"])
        self_check = "".join(f"<li>{esc(item)}</li>" for item in lesson["self_check"])
        example = lesson["worked_example"]
        lesson_sections.append(
            f"""
            <section class="lesson">
              <div class="eyebrow">Actual training module / Day {esc(lesson['day'])}</div>
              <h2>{esc(lesson['title'])}</h2>
              <h3>Objective</h3>
              <p>{esc(lesson['teaching_objective'])}</p>
              <h3>Plain-English Lesson</h3>
              {body}
              <h3>Key Terms</h3>
              <ul>{terms}</ul>
              <h3>Worked Example</h3>
              <p><strong>Scenario:</strong> {esc(example['scenario'])}</p>
              <p><strong>Weak version:</strong> {esc(example['bad_pitch'])}</p>
              <p><strong>Why weak:</strong> {esc(example['why_bad'])}</p>
              <p><strong>Better version:</strong> {esc(example['good_pitch'])}</p>
              <h3>Practice Steps</h3>
              <ul>{practice}</ul>
              <h3>Self-Check</h3>
              <ul>{self_check}</ul>
              <div class="answer"><strong>Answer key:</strong> {esc(lesson['answer_key'])}</div>
            </section>
            """
        )
    stop_lines = "".join(f"<li>{esc(item)}</li>" for item in packet["authority_boundary"])
    return f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>WF75 Academy Training Handout</title>
<style>
  @page {{ size: Letter; margin: 0.42in; }}
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; font-family: Arial, Helvetica, sans-serif; color: #162033; background: #ffffff; font-size: 11px; line-height: 1.38; }}
  .cover {{ min-height: 9.7in; padding: 0.35in 0 0; page-break-after: always; }}
  .kicker {{ color: #62758d; text-transform: uppercase; letter-spacing: .12em; font-size: 9px; font-weight: 700; }}
  h1 {{ color: #0e294a; font-size: 30px; margin: 0.08in 0 0.12in; line-height: 1.05; }}
  h2 {{ color: #0e294a; font-size: 18px; margin: 0 0 7px; }}
  h3 {{ color: #315d89; font-size: 10px; text-transform: uppercase; letter-spacing: .08em; margin: 9px 0 4px; }}
  p {{ margin: 4px 0 8px; }}
  ul {{ margin: 4px 0 8px 17px; padding: 0; }}
  li {{ margin: 3px 0; }}
  .hero {{ border-top: 5px solid #12345c; padding-top: 18px; }}
  .summary {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 8px; margin: 18px 0; }}
  .metric {{ border: 1px solid #d7e0ea; padding: 10px; min-height: 0.82in; }}
  .metric strong {{ display: block; font-size: 16px; color: #0e294a; margin-top: 4px; }}
  .card {{ border-top: 2px solid #12345c; padding: 12px 0; page-break-inside: avoid; }}
  .lesson {{ page-break-before: always; border-top: 5px solid #12345c; padding-top: 14px; }}
  .answer {{ border: 1px solid #b7c7da; background: #f4f8fd; padding: 9px; margin-top: 8px; }}
  table {{ width: 100%; border-collapse: collapse; table-layout: fixed; margin: 10px 0; }}
  th, td {{ border: 1px solid #d7e0ea; padding: 7px; vertical-align: top; }}
  th {{ background: #12345c; color: white; text-align: left; }}
  .warning {{ border: 1px solid #d5a736; background: #fff8df; padding: 10px; margin-top: 14px; }}
</style>
</head>
<body>
<section class="cover">
  <div class="hero">
    <div class="kicker">Internal training / fake-scenario practice only</div>
    <h1>WF75 Academy / Training Desk</h1>
    <p>{esc(packet['purpose'])}</p>
    <div class="summary">
      <div class="metric">Daily cadence<strong>{esc(packet['current_training_posture']['recommended_daily_time'])}</strong></div>
      <div class="metric">Weekly review<strong>{esc(packet['current_training_posture']['weekly_review_time'])}</strong></div>
      <div class="metric">First target<strong>SMB workflow packets</strong></div>
    </div>
    <h3>Repeatable session format</h3>
    <ul>{''.join(f"<li>{esc(item)}</li>" for item in packet['repeatable_session_format'])}</ul>
    <div class="warning"><strong>Boundary:</strong> This is internal readiness training. No real customer data, outreach, credentials, implementation, external delivery, spending, public launch, guaranteed ROI, legal/compliance/security readiness, or certification claim.</div>
  </div>
</section>
{''.join(lesson_sections)}
{''.join(session_cards)}
<section class="card">
  <h2>Tool Curriculum</h2>
  <table>
    <tr><th>Tool</th><th>Learn for</th><th>Starter pattern</th><th>Do not do yet</th></tr>
    {tool_rows}
  </table>
</section>
<section class="card">
  <h2>Stop Lines</h2>
  <ul>{stop_lines}</ul>
</section>
</body>
</html>"""


def render_simulation_html(packet: dict[str, Any]) -> str:
    sessions = packet["daily_micro_learning_plan"]
    lessons_by_day = {lesson["day"]: lesson for lesson in packet.get("lesson_modules", [])}
    slides = [
        {
            "title": "WF75 Academy Simulation Deck",
            "body": f"Updated {packet['generated_at_utc']}. This deck now starts each available lesson with actual training content, worked examples, and practice checks.",
            "prompt": "Press Next. Read the lesson slide, answer the worked-example prompt out loud, then bring your explain-back to chat for scoring.",
        }
    ]
    for session in sessions:
        lesson = lessons_by_day.get(session["day"])
        if lesson:
            slides.append(
                {
                    "title": f"Lesson: {lesson['title']}",
                    "body": " ".join(lesson["plain_english_lesson"][:2]),
                    "prompt": f"Practice: {' '.join(lesson['practice_steps'])}",
                }
            )
            example = lesson["worked_example"]
            slides.append(
                {
                    "title": f"Worked Example: Day {lesson['day']}",
                    "body": f"Weak: {example['bad_pitch']} Better: {example['good_pitch']}",
                    "prompt": f"Self-check: {' '.join(lesson['self_check'])}",
                }
            )
        slides.append(
            {
                "title": f"Day {session['day']}: {session['title']}",
                "body": f"Scenario: {session['fake_scenario']}",
                "prompt": f"Activity: {session['practice']} Acceptance: {session['acceptance']}",
            }
        )
    slides_json = json.dumps(slides)
    return f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>WF75 Academy Simulation Deck</title>
<style>
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; font-family: Arial, Helvetica, sans-serif; color: #132235; background: #f5f7fb; }}
  main {{ min-height: 100vh; display: grid; grid-template-rows: 1fr auto; }}
  .slide {{ padding: 8vh 8vw; display: grid; align-content: center; gap: 22px; }}
  .kicker {{ text-transform: uppercase; letter-spacing: .14em; color: #667b91; font-size: 13px; font-weight: 700; }}
  h1 {{ font-size: clamp(34px, 6vw, 72px); line-height: 1; margin: 0; max-width: 980px; }}
  p {{ font-size: clamp(20px, 2.4vw, 32px); line-height: 1.25; max-width: 940px; margin: 0; }}
  .prompt {{ border-left: 6px solid #12345c; background: white; padding: 20px 24px; font-size: clamp(18px, 2vw, 24px); max-width: 980px; }}
  .controls {{ display: flex; justify-content: space-between; align-items: center; padding: 18px 8vw; background: #0f294a; color: white; }}
  button {{ border: 1px solid rgba(255,255,255,.45); background: transparent; color: white; padding: 10px 16px; font-size: 16px; cursor: pointer; }}
  button:disabled {{ opacity: .4; cursor: default; }}
</style>
</head>
<body>
<main>
  <section class="slide">
    <div class="kicker" id="count"></div>
    <h1 id="title"></h1>
    <p id="body"></p>
    <div class="prompt" id="prompt"></div>
  </section>
  <nav class="controls">
    <button id="prev">Previous</button>
    <div>Internal practice only / fake scenarios only</div>
    <button id="next">Next</button>
  </nav>
</main>
<script>
const slides = {slides_json};
let i = 0;
const title = document.getElementById('title');
const body = document.getElementById('body');
const prompt = document.getElementById('prompt');
const count = document.getElementById('count');
const prev = document.getElementById('prev');
const next = document.getElementById('next');
function render() {{
  const s = slides[i];
  count.textContent = `Slide ${{i + 1}} of ${{slides.length}}`;
  title.textContent = s.title;
  body.textContent = s.body;
  prompt.textContent = s.prompt;
  prev.disabled = i === 0;
  next.disabled = i === slides.length - 1;
}}
prev.addEventListener('click', () => {{ if (i > 0) {{ i--; render(); }} }});
next.addEventListener('click', () => {{ if (i < slides.length - 1) {{ i++; render(); }} }});
document.addEventListener('keydown', (event) => {{
  if (event.key === 'ArrowRight' && i < slides.length - 1) {{ i++; render(); }}
  if (event.key === 'ArrowLeft' && i > 0) {{ i--; render(); }}
}});
render();
</script>
</body>
</html>"""


def render_pdf(html_path: Path, pdf_path: Path) -> dict[str, Any]:
    browser = find_browser()
    if not browser:
        return {"status": "pdf_unavailable", "reason": "no headless Edge/Chrome browser found"}
    cmd = [
        browser,
        "--headless=new",
        "--disable-gpu",
        "--no-first-run",
        "--disable-extensions",
        f"--print-to-pdf={pdf_path}",
        str(html_path),
    ]
    proc = subprocess.run(cmd, cwd=str(ROOT), text=True, capture_output=True, timeout=90)
    pdf_ok = proc.returncode == 0 and pdf_path.exists() and pdf_path.stat().st_size > 0
    return {
        "status": "ok" if pdf_ok else "pdf_failed",
        "browser": browser,
        "returncode": proc.returncode,
        "stdout_tail": proc.stdout[-1000:],
        "stderr_tail": proc.stderr[-1000:],
        "pdf_size_bytes": pdf_path.stat().st_size if pdf_path.exists() else 0,
    }


def render_pptx(packet: dict[str, Any], pptx_path: Path) -> dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        from pptx.util import Inches, Pt
    except Exception as exc:
        return {"status": "pptx_unavailable", "reason": str(exc)}

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    navy = RGBColor(18, 52, 92)
    text = RGBColor(22, 32, 51)
    muted = RGBColor(88, 105, 125)
    pale = RGBColor(246, 248, 252)

    def add_title(slide: Any, title: str, subtitle: str = "") -> None:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(0.42), Inches(11.9), Inches(0.8))
        frame = box.text_frame
        p = frame.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.size = Pt(30)
        run.font.bold = True
        run.font.color.rgb = text
        if subtitle:
            sub = slide.shapes.add_textbox(Inches(0.72), Inches(1.14), Inches(11.4), Inches(0.42))
            sub_frame = sub.text_frame
            sub_run = sub_frame.paragraphs[0].add_run()
            sub_run.text = subtitle
            sub_run.font.size = Pt(14)
            sub_run.font.color.rgb = muted

    def add_box(slide: Any, title: str, body: str, left: float, top: float, width: float, height: float) -> None:
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = pale
        shape.line.color.rgb = RGBColor(216, 225, 236)
        tx = slide.shapes.add_textbox(Inches(left + 0.18), Inches(top + 0.14), Inches(width - 0.36), Inches(height - 0.28))
        frame = tx.text_frame
        frame.word_wrap = True
        title_run = frame.paragraphs[0].add_run()
        title_run.text = title
        title_run.font.size = Pt(17)
        title_run.font.bold = True
        title_run.font.color.rgb = navy
        p = frame.add_paragraph()
        body_run = p.add_run()
        body_run.text = body
        body_run.font.size = Pt(14)
        body_run.font.color.rgb = text

    slide = prs.slides.add_slide(blank)
    add_title(slide, "WF75 Academy Activity Deck", "Internal training / fake-scenario practice only")
    add_box(slide, "Goal", packet["current_training_posture"]["first_skill_goal"], 0.8, 1.8, 5.7, 1.6)
    add_box(slide, "Cadence", f"Daily: {packet['current_training_posture']['recommended_daily_time']}\nWeekly review: {packet['current_training_posture']['weekly_review_time']}", 6.8, 1.8, 5.4, 1.6)
    add_box(slide, "Stop line", "No real customer data, outreach, credentials, implementation, external delivery, spending, public launch, or ROI/compliance claims.", 0.8, 4.0, 11.4, 1.35)

    for session in packet["daily_micro_learning_plan"]:
        slide = prs.slides.add_slide(blank)
        add_title(slide, f"Day {session['day']}: {session['title']}", f"{session['duration_minutes']} minutes")
        add_box(slide, "Scenario", session["fake_scenario"], 0.75, 1.55, 3.8, 1.45)
        add_box(slide, "Activity", session["practice"], 4.75, 1.55, 3.8, 1.45)
        add_box(slide, "Explain-back", "Randall states the workflow, bottleneck, next action, and stop line in plain English.", 8.75, 1.55, 3.8, 1.45)
        learn_text = "\n".join(f"- {item}" for item in session["learn"])
        add_box(slide, "Learning points", learn_text, 0.75, 3.45, 5.7, 2.1)
        add_box(slide, "Acceptance", session["acceptance"], 6.75, 3.45, 5.8, 2.1)

    for lesson in packet.get("lesson_modules", []):
        slide = prs.slides.add_slide(blank)
        add_title(slide, f"Training: {lesson['title']}", "Actual lesson content")
        add_box(slide, "Objective", lesson["teaching_objective"], 0.75, 1.4, 5.8, 1.2)
        add_box(slide, "Core lesson", "\n".join(lesson["plain_english_lesson"][:3]), 0.75, 2.9, 5.8, 2.6)
        terms = "\n".join(f"{item['term']}: {item['definition']}" for item in lesson["key_terms"][:3])
        add_box(slide, "Key terms", terms, 6.85, 1.4, 5.75, 2.0)
        add_box(slide, "Practice", "\n".join(lesson["practice_steps"]), 6.85, 3.75, 5.75, 1.75)

        example = lesson["worked_example"]
        slide = prs.slides.add_slide(blank)
        add_title(slide, f"Worked Example: {lesson['title']}", "Compare weak vs better framing")
        add_box(slide, "Scenario", example["scenario"], 0.75, 1.35, 11.8, 1.0)
        add_box(slide, "Weak version", example["bad_pitch"], 0.75, 2.75, 5.7, 1.4)
        add_box(slide, "Why weak", example["why_bad"], 6.85, 2.75, 5.7, 1.4)
        add_box(slide, "Better version", example["good_pitch"], 0.75, 4.55, 11.8, 1.35)

        slide = prs.slides.add_slide(blank)
        add_title(slide, f"Practice Check: {lesson['title']}", "Use this slide for feedback")
        add_box(slide, "Self-check", "\n".join(lesson["self_check"]), 0.75, 1.45, 5.75, 3.0)
        add_box(slide, "Answer key", lesson["answer_key"], 6.85, 1.45, 5.7, 3.0)
        add_box(slide, "Your turn", "Explain it back in 90 seconds. Veritas scores clarity, scope control, workflow thinking, tool judgment, safety, and explain-back.", 0.75, 5.0, 11.8, 1.0)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Tool Pattern Decision Drill", "Pick the simplest pattern before recommending a platform")
    for idx, tool in enumerate(packet["tool_curriculum"][:5]):
        left = 0.75 + (idx % 2) * 6.05
        top = 1.45 + (idx // 2) * 1.7
        add_box(slide, tool["tool"], f"{tool['learn_for']}\nStarter: {tool['starter_pattern']}", left, top, 5.65, 1.35)

    prs.save(pptx_path)
    return {"status": "ok", "pptx_size_bytes": pptx_path.stat().st_size if pptx_path.exists() else 0}


def render_docx_book(packet: dict[str, Any], docx_path: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except Exception as exc:
        return {"status": "docx_unavailable", "reason": str(exc)}

    doc = Document()
    doc.add_heading("WF75 Academy Book", 0)
    doc.add_paragraph("Internal training reference for Randall. This book is cumulative review material, not a certification or customer-delivery artifact.")
    doc.add_paragraph(f"Generated: {packet['generated_at_utc']}")
    doc.add_heading("How To Use This Book", level=1)
    for item in [
        "Read the lesson before practice.",
        "Use the worked example to understand good vs weak framing.",
        "Complete the practice steps out loud or in chat.",
        "Use the self-check before asking Veritas for scoring.",
        "Return to this book for review before fake scenarios and later controlled pilots.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Training Path", level=1)
    for session in packet["daily_micro_learning_plan"]:
        doc.add_paragraph(f"Day {session['day']}: {session['title']} ({session['duration_minutes']} minutes)", style="List Bullet")

    doc.add_heading("Lesson Modules", level=1)
    for lesson in packet.get("lesson_modules", []):
        doc.add_heading(f"Day {lesson['day']}: {lesson['title']}", level=2)
        doc.add_paragraph(f"Objective: {lesson['teaching_objective']}")
        doc.add_heading("Plain-English Lesson", level=3)
        for paragraph in lesson["plain_english_lesson"]:
            doc.add_paragraph(paragraph)
        doc.add_heading("Key Terms", level=3)
        for item in lesson["key_terms"]:
            doc.add_paragraph(f"{item['term']}: {item['definition']}", style="List Bullet")
        example = lesson["worked_example"]
        doc.add_heading("Worked Example", level=3)
        doc.add_paragraph(f"Scenario: {example['scenario']}")
        doc.add_paragraph(f"Weak version: {example['bad_pitch']}")
        doc.add_paragraph(f"Why weak: {example['why_bad']}")
        doc.add_paragraph(f"Better version: {example['good_pitch']}")
        doc.add_heading("Practice Steps", level=3)
        for item in lesson["practice_steps"]:
            doc.add_paragraph(item, style="List Number")
        doc.add_heading("Self-Check", level=3)
        for item in lesson["self_check"]:
            doc.add_paragraph(item, style="List Bullet")
        doc.add_heading("Answer Key", level=3)
        doc.add_paragraph(lesson["answer_key"])

    doc.add_heading("Tool Curriculum", level=1)
    for tool in packet["tool_curriculum"]:
        doc.add_heading(tool["tool"], level=2)
        doc.add_paragraph(f"Learn for: {tool['learn_for']}")
        doc.add_paragraph(f"Starter pattern: {tool['starter_pattern']}")
        doc.add_paragraph(f"Do not do yet: {tool['do_not_do_yet']}")

    doc.add_heading("Stop Lines", level=1)
    for item in packet["authority_boundary"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(docx_path)
    return {"status": "ok", "docx_size_bytes": docx_path.stat().st_size if docx_path.exists() else 0}


def write_training_assets(packet: dict[str, Any]) -> dict[str, Any]:
    ACADEMY.mkdir(parents=True, exist_ok=True)
    atomic_write_json(TRAINING_JSON, packet)
    atomic_write_text(TRAINING_MD, render_markdown(packet))
    atomic_write_text(TRAINING_HTML, render_html_handout(packet))
    atomic_write_text(TRAINING_SIM_HTML, render_simulation_html(packet))
    pdf_result = render_pdf(TRAINING_HTML, TRAINING_PDF)
    docx_result = render_docx_book(packet, TRAINING_DOCX)
    pptx_result = render_pptx(packet, TRAINING_PPTX)
    manifest = {
        "schema": "veritas.wf75_academy_manifest.v1",
        "generated_at_utc": utc_now(),
        "status": "ready" if pdf_result.get("status") == "ok" and pptx_result.get("status") == "ok" and docx_result.get("status") == "ok" else "warning",
        "outputs": {
            "json": rel(TRAINING_JSON),
            "markdown": rel(TRAINING_MD),
            "html_handout": rel(TRAINING_HTML),
            "pdf_handout": rel(TRAINING_PDF),
            "docx_book": rel(TRAINING_DOCX),
            "pptx_activity_deck": rel(TRAINING_PPTX),
            "html_simulation_deck": rel(TRAINING_SIM_HTML),
            "manifest": rel(TRAINING_MANIFEST),
        },
        "pdf_result": pdf_result,
        "docx_result": docx_result,
        "pptx_result": pptx_result,
        "authority_boundary": AUTHORITY_FALSE,
        "training_posture": {
            "adult_learning": True,
            "uses_activities": True,
            "uses_simulations": True,
            "reading_only": False,
            "future_app_candidate": True,
        },
        "validation": {"status": "pending", "errors": [], "warnings": []},
    }
    errors: list[str] = []
    warnings: list[str] = []
    for key, value in manifest["outputs"].items():
        path = ROOT / value
        if key != "manifest" and (not path.exists() or path.stat().st_size <= 0):
            errors.append(f"missing_or_empty_output:{value}")
    if pdf_result.get("status") != "ok":
        warnings.append(f"pdf_not_ready:{pdf_result.get('status')}")
    if docx_result.get("status") != "ok":
        warnings.append(f"docx_not_ready:{docx_result.get('status')}")
    if pptx_result.get("status") != "ok":
        warnings.append(f"pptx_not_ready:{pptx_result.get('status')}")
    manifest["validation"] = {"status": "ok" if not errors else "blocked", "errors": errors, "warnings": warnings}
    if errors:
        manifest["status"] = "blocked"
    atomic_write_json(TRAINING_MANIFEST, manifest)
    return manifest


def validate(packet: dict[str, Any]) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    if packet.get("authority_boundary") != AUTHORITY_FALSE:
        errors.append("authority_boundary_changed")
    if len(packet.get("daily_micro_learning_plan", [])) < 7:
        errors.append("daily_plan_below_7_sessions")
    if not packet.get("tool_curriculum"):
        errors.append("tool_curriculum_missing")
    if not packet.get("lesson_modules"):
        errors.append("lesson_modules_missing")
    if packet.get("current_training_posture", {}).get("academy_live") is not True:
        errors.append("academy_not_live")
    if packet.get("authority_boundary", {}).get("external_delivery_allowed") is not False:
        errors.append("external_delivery_boundary_not_false")
    return {"status": "ok" if not errors else "blocked", "errors": errors, "warnings": warnings}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build WF75 Academy / Training Desk artifacts.")
    parser.add_argument("--write", action="store_true")
    parser.add_argument("--write-md", action="store_true")
    parser.add_argument("--write-training-assets", action="store_true")
    parser.add_argument("--validate", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    packet = build_packet()
    packet["validation"] = validate(packet)
    if args.write:
        atomic_write_json(OUT_JSON, packet)
    if args.write_md:
        OUT_MD.write_text(render_markdown(packet), encoding="utf-8")
    training_manifest = write_training_assets(packet) if args.write_training_assets else None
    result = {
        "status": packet["validation"]["status"],
        "outputs": {
            "json": rel(OUT_JSON),
            "markdown": rel(OUT_MD) if args.write_md else None,
            "training_manifest": rel(TRAINING_MANIFEST) if training_manifest else None,
            "training_pdf": rel(TRAINING_PDF) if training_manifest else None,
            "training_docx": rel(TRAINING_DOCX) if training_manifest else None,
            "training_pptx": rel(TRAINING_PPTX) if training_manifest else None,
            "training_simulation": rel(TRAINING_SIM_HTML) if training_manifest else None,
        },
        "daily_sessions": len(packet["daily_micro_learning_plan"]),
        "tool_modules": len(packet["tool_curriculum"]),
        "training_asset_status": training_manifest.get("status") if training_manifest else None,
        "validation": packet["validation"],
    }
    if training_manifest:
        result["training_asset_validation"] = training_manifest.get("validation")
    print(json.dumps(result, indent=2))
    if args.validate and packet["validation"]["status"] != "ok":
        return 2
    if args.validate and training_manifest and training_manifest.get("validation", {}).get("status") != "ok":
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
