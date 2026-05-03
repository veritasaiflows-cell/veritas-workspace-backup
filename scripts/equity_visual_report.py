import argparse
import json
from pathlib import Path

import yfinance as yf
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

WORKSPACE = Path(__file__).resolve().parents[1]
TMP = WORKSPACE / "tmp"
PLAYBOOKS = WORKSPACE / "06. Playbooks"
TMP.mkdir(parents=True, exist_ok=True)
PLAYBOOKS.mkdir(parents=True, exist_ok=True)

BG = "#f5f1e8"
CARD = "#ede7da"
TEXT = "#111111"
MUTED = "#3f3f3f"
GREEN = "#0d8a53"
RED = "#b94728"
BLUE1 = "#9fc0e6"
BLUE2 = "#2f79c8"
BLUE3 = "#205e9e"
BLUE4 = "#163f6b"


def money(v):
    if v is None:
        return "N/A"
    a = abs(v)
    if a >= 1_000_000_000:
        return f"${v/1_000_000_000:.1f}B"
    if a >= 1_000_000:
        return f"${v/1_000_000:.1f}M"
    return f"${v:,.0f}"


def pct_ratio(v, digits=1):
    if v is None:
        return "N/A"
    value = float(v)
    if abs(value) > 1:
        return f"{value:.{digits}f}%"
    return f"{value*100:.{digits}f}%"


def pct_num(v, digits=1):
    if v is None:
        return "N/A"
    return f"{v:.{digits}f}%"


def load_fonts():
    candidates = [
        r"C:\Windows\Fonts\aptos.ttf",
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\segoeui.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
    ]
    path = next((p for p in candidates if Path(p).exists()), None)
    if path:
        return {
            "title": ImageFont.truetype(path, 34),
            "h1": ImageFont.truetype(path, 26),
            "h2": ImageFont.truetype(path, 21),
            "body": ImageFont.truetype(path, 17),
            "small": ImageFont.truetype(path, 14),
            "metric": ImageFont.truetype(path, 30),
        }
    d = ImageFont.load_default()
    return {k: d for k in ["title", "h1", "h2", "body", "small", "metric"]}


def round_rect(draw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def _extract_income_history(ticker):
    income = ticker.income_stmt
    if income is None or income.empty:
        return {
            "eps_labels": ["N/A"],
            "eps_values": [0],
            "rev_labels": ["N/A"],
            "rev_values": [0],
        }

    cols = list(income.columns)
    cols_sorted = sorted(cols)
    actual_cols = cols_sorted[-3:] if len(cols_sorted) >= 3 else cols_sorted

    eps_values = []
    rev_values = []
    labels = []
    for col in actual_cols:
        year = str(getattr(col, "year", str(col)[:4]))
        labels.append(year)
        eps = None
        rev = None
        if "Diluted EPS" in income.index:
            try:
                eps = float(income.loc["Diluted EPS", col])
            except Exception:
                eps = None
        if "Total Revenue" in income.index:
            try:
                rev = float(income.loc["Total Revenue", col]) / 1_000_000_000
            except Exception:
                rev = None
        eps_values.append(round(eps or 0, 2))
        rev_values.append(round(rev or 0, 1))

    info = ticker.info
    cal = ticker.calendar or {}
    next_year = None
    if labels:
        try:
            next_year = str(int(labels[-1]) + 1) + "E"
        except Exception:
            next_year = None
    est_eps = info.get("epsCurrentYear")
    rev_avg = cal.get("Revenue Average") if isinstance(cal, dict) else None
    if next_year and est_eps:
        labels.append(next_year)
        eps_values.append(round(float(est_eps), 2))
        if rev_avg:
            rev_values.append(round(float(rev_avg) * 4 / 1_000_000_000, 1))
        elif rev_values:
            rev_values.append(rev_values[-1])
        else:
            rev_values.append(0)

    return {
        "eps_labels": labels,
        "eps_values": eps_values,
        "rev_labels": labels,
        "rev_values": rev_values,
    }


def _derive_verdict(latest, target_upside):
    close = latest["close"]
    ma20 = latest["ma20"]
    ma50 = latest["ma50"]
    ma200 = latest["ma200"]
    near_high = close >= latest["high_52w"] * 0.98
    near_resistance = close >= latest["recent_resistance"] * 0.99

    if close < ma20 and close < ma50 and close < ma200:
        return "DO NOT DEPLOY", "Price is below the 20-day, 50-day, and 200-day moving averages. The chart needs real repair before a fresh long case is actionable."
    if close > ma20 > ma50 > ma200:
        if near_high or near_resistance:
            return "PULLBACK ONLY", "Trend quality is strong, but price is pressing highs and current entry quality is poor. Wait for a disciplined pullback or a better-defined reset."
        return "CONSTRUCTIVE / WATCH FOR ENTRY", "Trend quality is constructive and the chart is healthier, but entry discipline still matters."
    if close > ma50 and close > ma200:
        return "WATCH / CONSTRUCTIVE", "The structure is improving, but it is not yet strong enough to treat as a clean deployment setup."
    if close > ma200:
        return "WATCH / MIXED", "Longer-term trend support survives, but the intermediate setup is mixed and needs cleaner confirmation."
    if target_upside is not None and target_upside < 0:
        return "WATCH / BENCH", "Analyst upside is not compelling at the current price, and the setup does not justify forcing capital."
    return "WATCH / BENCH", "Technically mixed or underdefined despite acceptable business quality."


def fetch_base_data(ticker_symbol):
    ticker = yf.Ticker(ticker_symbol)
    hist = ticker.history(period="3y", interval="1d", auto_adjust=False).dropna()
    info = ticker.info
    close = hist["Close"]
    ma20 = close.rolling(20).mean()
    ma50 = close.rolling(50).mean()
    ma200 = close.rolling(200).mean()
    ytd_base = close[close.index.year == close.index[-1].year].iloc[0]

    latest = {
        "date": str(hist.index[-1].date()),
        "close": float(close.iloc[-1]),
        "open": float(hist["Open"].iloc[-1]),
        "high": float(hist["High"].iloc[-1]),
        "low": float(hist["Low"].iloc[-1]),
        "volume": int(hist["Volume"].iloc[-1]),
        "ma20": float(ma20.iloc[-1]),
        "ma50": float(ma50.iloc[-1]),
        "ma200": float(ma200.iloc[-1]),
        "high_52w": float(hist["High"].tail(252).max()),
        "low_52w": float(hist["Low"].tail(252).min()),
        "recent_support": float(close.tail(20).min()),
        "recent_resistance": float(close.tail(20).max()),
        "one_year_return_pct": float((close.iloc[-1] / close.iloc[-252] - 1) * 100),
        "ytd_return_pct": float((close.iloc[-1] / ytd_base - 1) * 100),
    }

    target_upside = None
    if info.get("targetMeanPrice"):
        target_upside = ((info.get("targetMeanPrice") or 0) / (info.get("currentPrice") or latest["close"]) - 1) * 100
    off_peak = (latest["close"] / latest["high_52w"] - 1) * 100

    snapshot = {
        "last_close": f"${latest['close']:.2f}",
        "consensus_target": f"${(info.get('targetMeanPrice') or 0):.2f}" if info.get('targetMeanPrice') else "N/A",
        "market_cap": money(info.get("marketCap")),
        "range_52w": f"${latest['low_52w']:.0f}-${latest['high_52w']:.0f}",
        "dividend_yield": pct_ratio(info.get("dividendYield")),
        "forward_pe": f"~{info.get('forwardPE'):.0f}x" if info.get("forwardPE") else "N/A",
        "target_upside": pct_num(target_upside),
        "off_peak": pct_num(off_peak),
        "inst_owned": pct_ratio(info.get("heldPercentInstitutions")),
    }

    fundamentals = {
        "company": info.get("longName") or ticker_symbol,
        "exchange": info.get("fullExchangeName") or info.get("exchange") or "N/A",
        "sector": info.get("sectorDisp") or info.get("sector") or "N/A",
        "industry": info.get("industryDisp") or info.get("industry") or "N/A",
        "employees": info.get("fullTimeEmployees"),
        "market_cap": info.get("marketCap"),
        "enterprise_value": info.get("enterpriseValue"),
        "trailing_pe": info.get("trailingPE"),
        "forward_pe": info.get("forwardPE"),
        "price_to_sales": info.get("priceToSalesTrailing12Months"),
        "price_to_book": info.get("priceToBook"),
        "dividend_yield": info.get("dividendYield"),
        "profit_margin": info.get("profitMargins"),
        "operating_margin": info.get("operatingMargins"),
        "gross_margin": info.get("grossMargins"),
        "roe": info.get("returnOnEquity"),
        "roa": info.get("returnOnAssets"),
        "revenue_growth": info.get("revenueGrowth"),
        "earnings_growth": info.get("earningsGrowth"),
        "free_cash_flow": info.get("freeCashflow"),
        "operating_cash_flow": info.get("operatingCashflow"),
        "total_revenue": info.get("totalRevenue"),
        "total_debt": info.get("totalDebt"),
        "total_cash": info.get("totalCash"),
        "current_ratio": info.get("currentRatio"),
        "quick_ratio": info.get("quickRatio"),
        "debt_to_equity": info.get("debtToEquity"),
        "analyst_rating": info.get("averageAnalystRating"),
        "recommendation_key": info.get("recommendationKey"),
        "analyst_count": info.get("numberOfAnalystOpinions"),
        "target_mean_price": info.get("targetMeanPrice"),
        "target_median_price": info.get("targetMedianPrice"),
        "target_low_price": info.get("targetLowPrice"),
        "target_high_price": info.get("targetHighPrice"),
        "beta": info.get("beta"),
        "held_percent_institutions": info.get("heldPercentInstitutions"),
    }

    verdict, verdict_reason = _derive_verdict(latest, target_upside)

    return {
        "ticker": ticker_symbol,
        "latest": latest,
        "snapshot": snapshot,
        "fundamentals": fundamentals,
        "hist_fin": _extract_income_history(ticker),
        "verdict": verdict,
        "verdict_reason": verdict_reason,
    }


def draw_price_panel(data, out_path, fonts):
    img = Image.new("RGB", (1700, 860), BG)
    d = ImageDraw.Draw(img)
    d.text((60, 26), f"{data['ticker']} PRICE SNAPSHOT", font=fonts["title"], fill=TEXT)
    cards = [
        ("Last close", data["snapshot"]["last_close"], f"52-wk high: ${data['latest']['high_52w']:.2f}", RED),
        ("Consensus target", data["snapshot"]["consensus_target"], f"{data['snapshot']['target_upside']} implied upside", GREEN if "-" not in data['snapshot']['target_upside'] else RED),
        ("Market cap", data["snapshot"]["market_cap"], f"{data['snapshot']['inst_owned']} institutional owned", MUTED),
        ("52-wk range", data["snapshot"]["range_52w"], f"{data['snapshot']['off_peak']} from peak", RED if "-" in data['snapshot']['off_peak'] else GREEN),
        ("Dividend yield", data["snapshot"]["dividend_yield"], "Market data field", MUTED),
        ("P/E (fwd)", data["snapshot"]["forward_pe"], "Forward earnings multiple", MUTED),
    ]
    x_positions = [60, 380, 700, 1020, 1340]
    y1 = 100
    w, h = 280, 180
    for i, (title, value, note, note_color) in enumerate(cards[:5]):
        round_rect(d, (x_positions[i], y1, x_positions[i] + w, y1 + h), 18, CARD)
        d.text((x_positions[i] + 24, y1 + 20), title, font=fonts["h2"], fill=TEXT)
        d.text((x_positions[i] + 24, y1 + 68), value, font=fonts["metric"], fill=TEXT)
        d.text((x_positions[i] + 24, y1 + 128), note, font=fonts["body"], fill=note_color)
    round_rect(d, (60, 310, 340, 490), 18, CARD)
    d.text((84, 330), cards[5][0], font=fonts["h2"], fill=TEXT)
    d.text((84, 378), cards[5][1], font=fonts["metric"], fill=TEXT)
    d.text((84, 438), cards[5][2], font=fonts["body"], fill=MUTED)
    img.save(out_path)


def draw_bar_chart(draw, fonts, origin, size, labels, values, colors, title, y_prefix="", y_suffix=""):
    x0, y0 = origin
    w, h = size
    draw.text((x0, y0 - 42), title, font=fonts["h1"], fill=TEXT)
    plot_top = y0 + 10
    plot_bottom = y0 + h
    plot_left = x0 + 40
    plot_right = x0 + w
    maxv = max(values) if values else 0
    maxv = maxv * 1.15 if maxv > 0 else 1
    for i in range(0, 6):
        y = plot_bottom - int((plot_bottom - plot_top) * i / 5)
        draw.line((plot_left, y, plot_right, y), fill="#d2cbc0", width=2)
        label = f"{y_prefix}{maxv*i/5:.0f}{y_suffix}"
        draw.text((x0, y - 8), label, font=fonts["small"], fill=MUTED)
    n = len(values)
    gap = 44
    bar_w = int((plot_right - plot_left - gap * max(n - 1, 0)) / max(n, 1))
    for i, (lab, val, col) in enumerate(zip(labels, values, colors[:n])):
        bx = plot_left + i * (bar_w + gap)
        bh = int((val / maxv) * (plot_bottom - plot_top - 10)) if maxv else 0
        by = plot_bottom - bh
        draw.rectangle((bx, by, bx + bar_w, plot_bottom), fill=col)
        tw = draw.textbbox((0, 0), lab, font=fonts["body"])[2]
        draw.text((bx + (bar_w - tw) / 2, plot_bottom + 14), lab, font=fonts["body"], fill=TEXT)


def draw_history_panel(data, out_path, fonts):
    img = Image.new("RGB", (1700, 760), BG)
    d = ImageDraw.Draw(img)
    colors = [BLUE1, BLUE2, BLUE3, BLUE4, GREEN]
    draw_bar_chart(d, fonts, (40, 70), (740, 560), data["hist_fin"]["eps_labels"], data["hist_fin"]["eps_values"], colors, "HISTORICAL EPS TRAJECTORY", y_prefix="$", y_suffix="")
    draw_bar_chart(d, fonts, (880, 70), (740, 560), data["hist_fin"]["rev_labels"], data["hist_fin"]["rev_values"], colors, "HISTORICAL REVENUE TRAJECTORY", y_prefix="$", y_suffix="B")
    img.save(out_path)


def set_normal_font(document):
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)


def add_title(document, title, subtitle):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(18)
    p2 = document.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.italic = True
    r2.font.size = Pt(10)


def add_metric_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Value"
    for key, value in rows:
        c = table.add_row().cells
        c[0].text = str(key)
        c[1].text = str(value)
    return table


def build_report(data, out_path, price_panel, history_panel, subtitle=None):
    latest = data["latest"]
    f = data["fundamentals"]
    doc = Document()
    set_normal_font(doc)
    add_title(doc, f"{f['company']} ({data['ticker']}) Visual Equity Report", subtitle or "Decision-grade report core updated with stronger executive summary, company snapshot, and real ticker-level history.")

    alert = doc.add_paragraph()
    alert_run = alert.add_run(
        f"VERITAS ALERT: {data['ticker']} closed at ${latest['close']:.2f} on {latest['date']}. "
        f"20-day / 50-day / 200-day levels are ${latest['ma20']:.2f} / ${latest['ma50']:.2f} / ${latest['ma200']:.2f}. "
        f"Current status: {data['verdict']}. {data['verdict_reason']}"
    )
    alert_run.bold = True

    doc.add_heading("1. Executive Summary", level=1)
    forward_pe_text = f"{f['forward_pe']:.1f}" if f['forward_pe'] else "N/A"
    doc.add_paragraph(
        f"{f['company']} is a real institutional-quality company in {f['sector']} / {f['industry']}, with market capitalization of {money(f['market_cap'])}, forward earnings multiple of {forward_pe_text}, and analyst support that still leans constructive. "
        f"The name remains worth covering even when the current setup is not clean enough to buy."
    )
    doc.add_paragraph(
        f"At the market level, {data['ticker']} closed at ${latest['close']:.2f}, with recent support near ${latest['recent_support']:.2f}, resistance near ${latest['recent_resistance']:.2f}, and a 52-week range of ${latest['low_52w']:.2f} to ${latest['high_52w']:.2f}. "
        f"Consensus target sits at {data['snapshot']['consensus_target']}, implying {data['snapshot']['target_upside']} from the current price."
    )
    doc.add_paragraph(
        f"The right current interpretation is {data['verdict']}. {data['verdict_reason']} "
        f"This report therefore emphasizes a clean company snapshot, visual price context, real annual financial history, and a simple action view rather than padding the output with generic commentary."
    )

    doc.add_heading("2. Company Snapshot", level=1)
    add_metric_table(doc, [
        ("Ticker", f"{data['ticker']} ({f['exchange']})"),
        ("Full name", f['company']),
        ("Sector / industry", f"{f['sector']} / {f['industry']}"),
        ("Market cap", money(f['market_cap'])),
        ("Enterprise value", money(f['enterprise_value'])),
        ("Employees", f"{f['employees']:,}" if f['employees'] else "N/A"),
        ("Dividend yield", pct_ratio(f['dividend_yield'])),
        ("Analyst rating", f['analyst_rating']),
        ("Consensus target", f"${f['target_mean_price']:.2f}" if f['target_mean_price'] else "N/A"),
        ("Current status", data['verdict']),
    ])

    doc.add_heading("3. Visual Snapshot", level=1)
    doc.add_picture(str(price_panel), width=Inches(6.7))

    doc.add_heading("4. Historical Trend", level=1)
    doc.add_picture(str(history_panel), width=Inches(6.7))

    doc.add_heading("5. Valuation and Quality Metrics", level=1)
    add_metric_table(doc, [
        ("Revenue", money(f['total_revenue'])),
        ("Free cash flow", money(f['free_cash_flow'])),
        ("Operating cash flow", money(f['operating_cash_flow'])),
        ("Trailing P/E", f"{f['trailing_pe']:.1f}" if f['trailing_pe'] else "N/A"),
        ("Forward P/E", f"{f['forward_pe']:.1f}" if f['forward_pe'] else "N/A"),
        ("Price / Sales", f"{f['price_to_sales']:.2f}" if f['price_to_sales'] else "N/A"),
        ("Price / Book", f"{f['price_to_book']:.2f}" if f['price_to_book'] else "N/A"),
        ("Profit margin", pct_ratio(f['profit_margin'])),
        ("Operating margin", pct_ratio(f['operating_margin'])),
        ("ROE", pct_ratio(f['roe'])),
        ("Current ratio", f"{f['current_ratio']:.2f}" if f['current_ratio'] else "N/A"),
        ("Quick ratio", f"{f['quick_ratio']:.2f}" if f['quick_ratio'] else "N/A"),
    ])

    doc.add_heading("6. Analyst Context", level=1)
    add_metric_table(doc, [
        ("Consensus target", f"${f['target_mean_price']:.2f}" if f['target_mean_price'] else "N/A"),
        ("Median target", f"${f['target_median_price']:.2f}" if f['target_median_price'] else "N/A"),
        ("Target range", f"${f['target_low_price']:.2f} to ${f['target_high_price']:.2f}" if f['target_low_price'] and f['target_high_price'] else "N/A"),
        ("Rating", f['analyst_rating']),
        ("Recommendation", str(f['recommendation_key']).title() if f['recommendation_key'] else "N/A"),
        ("Analyst count", f['analyst_count']),
    ])

    doc.add_heading("7. Action View", level=1)
    doc.add_paragraph(
        f"Verdict: {data['verdict']}. {data['verdict_reason']} "
        f"A future upgrade would require better trend posture, improved entry quality, and a cleaner risk-defined setup than the market currently offers."
    )

    doc.save(out_path)


def main():
    parser = argparse.ArgumentParser(description="Generate a reusable visual equity Word report.")
    parser.add_argument("ticker", help="Ticker symbol, e.g. RTX")
    parser.add_argument("--title-subtitle", default=None, help="Optional subtitle override")
    parser.add_argument("--out", default=None, help="Optional output path for the .docx")
    args = parser.parse_args()

    ticker = args.ticker.upper()
    data = fetch_base_data(ticker)
    fonts = load_fonts()

    safe = ticker.lower().replace(".", "-")
    price_panel = TMP / f"{safe}-price-panel.png"
    history_panel = TMP / f"{safe}-history-panel.png"
    data_json = TMP / f"{safe}-visual-report-data.json"
    out_path = Path(args.out) if args.out else PLAYBOOKS / f"{ticker} Visual Report - {data['latest']['date']}.docx"

    draw_price_panel(data, price_panel, fonts)
    draw_history_panel(data, history_panel, fonts)
    build_report(data, out_path, price_panel, history_panel, subtitle=args.title_subtitle)
    data_json.write_text(json.dumps(data, indent=2), encoding="utf-8")

    print(json.dumps({
        "ticker": ticker,
        "report_path": str(out_path),
        "price_panel": str(price_panel),
        "history_panel": str(history_panel),
        "data_json": str(data_json)
    }, indent=2))


if __name__ == "__main__":
    main()
